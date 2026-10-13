"""Gera o Relatorio de Faturamento em PowerPoint (.pptx), no modelo da Conexao (13,33 x 7,5 pol)."""
from __future__ import annotations

import io
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ASSETS = Path(__file__).parent / 'assets'
AZUL_ESCURO, AZUL_MEDIO, AZUL_CLARO, CINZA_LINHA, CINZA_TEXTO = '1F4E79', '2E75B6', '008EC9', 'F4F6F8', '454C52'
TONS_TXT = {'pos': '1B7F3B', 'neg': 'C62828', 'zero': '6B7280'}
CTX = {'gerado': '', 'pag': 0, 'total': 0}
TONS_CARD = {'azul': '1F4E79', 'laranja': 'B7791F', 'vermelho': 'C62828', 'verde': '1B7F3B'}


def _rgb(h):
    return RGBColor.from_string(h)


def _texto(slide, x, y, w, h, txt, tam, cor, negrito=False, italico=False, alinhar=PP_ALIGN.LEFT, ancora=MSO_ANCHOR.TOP, espaco=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ancora
    p = tf.paragraphs[0]
    p.alignment = alinhar
    r = p.add_run()
    r.text = txt
    f = r.font
    f.name, f.size, f.bold, f.italic = 'Calibri', Pt(tam), negrito, italico
    f.color.rgb = _rgb(cor)
    if espaco:
        r._r.get_or_add_rPr().set('spc', str(espaco))
    return tb


def _retangulo(slide, x, y, w, h, cor, forma=MSO_SHAPE.RECTANGLE, borda=None):
    s = slide.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(cor)
    if borda:
        s.line.color.rgb = _rgb(borda)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def _rodape(s):
    """Faixa azul com o site e, a direita, o titulo do relatorio, a data/hora em que foi gerado e a pagina."""
    _retangulo(s, 0, 7.29, 13.333, 0.21, AZUL_CLARO)
    _texto(s, 0.4, 7.32, 4.0, 0.15, 'www.gestaoconexao.com.br', 9, 'FFFFFF')
    _texto(s, 5.4, 7.32, 7.55, 0.15, f'Relatório de Faturamento  ·  Gerado em {CTX["gerado"]}  ·  {CTX["pag"]}/{CTX["total"]}', 9, 'FFFFFF', alinhar=PP_ALIGN.RIGHT)


def _base_conteudo(prs, secao, titulo, nota):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    CTX['pag'] += 1
    s.shapes.add_picture(str(ASSETS / 'conteudo_fundo.jpg'), 0, 0, prs.slide_width, prs.slide_height)
    s.shapes.add_picture(str(ASSETS / 'logo_colorida.png'), Inches(10.55), Inches(0.18), Inches(1.88), Inches(0.68))      # logo da Conexao em toda pagina
    _texto(s, 0.5, 0.22, 9.5, 0.25, secao, 13, AZUL_CLARO, negrito=True)
    _texto(s, 0.5, 0.5, 9.9, 0.45, titulo, 23 if len(titulo) <= 52 else (20 if len(titulo) <= 60 else 18), AZUL_ESCURO, negrito=True)
    _rodape(s)
    if nota:
        _texto(s, 0.35, 6.98, 12.6, 0.2, nota, 9, '6B7280', italico=True)
    return s


def _capa(prs, p):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(ASSETS / 'capa_fundo.jpg'), 0, 0, prs.slide_width, prs.slide_height)
    # a logo branca da Conexao ja faz parte da imagem de fundo da capa (no mesmo lugar do modelo)
    _texto(s, 0.4, 1.32, 11.5, 0.9, p['titulo'].upper(), 44, 'FFFFFF', negrito=True, espaco=300)
    ln = s.shapes.add_connector(1, Inches(0), Inches(2.23), Inches(3.41), Inches(2.23))
    ln.line.color.rgb = _rgb(AZUL_CLARO)
    ln.line.width = Pt(1.5)
    _texto(s, 0.4, 2.5, 11.5, 0.5, p['sub'].upper(), 24, 'FFFFFF', negrito=True, espaco=150)
    CTX['pag'] += 1
    _rodape(s)


def _celula(cel, valor, tam, cor, fundo, negrito=False, alinhar=PP_ALIGN.RIGHT):
    txt, tom = (valor if isinstance(valor, tuple) else (valor, None))
    cel.fill.solid()
    cel.fill.fore_color.rgb = _rgb(fundo)
    cel.margin_left = cel.margin_right = Emu(50800)
    cel.margin_top = cel.margin_bottom = Emu(12700)
    cel.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cel.text_frame
    tf.text = ''
    p = tf.paragraphs[0]
    p.alignment = alinhar
    r = p.add_run()
    r.text = str(txt)
    r.font.name, r.font.size, r.font.bold = 'Calibri', Pt(tam), negrito or bool(tom and tom != 'zero')
    r.font.color.rgb = _rgb(TONS_TXT[tom] if tom and fundo not in (AZUL_ESCURO, AZUL_MEDIO) else cor)


def _tabela(prs, p, s=None):
    if s is None:
        s = _base_conteudo(prs, p['secao'], p['titulo'], p.get('nota'))
    linhas, total, grupos, cab, larg = p['linhas'], p.get('total'), p['grupos'], p['cab'], p['larguras']
    n_lin = 2 + len(linhas) + (1 if total else 0)
    disponivel = 5.65
    alt = min(0.5, disponivel / n_lin)
    tam = p.get('tam') or (12 if alt >= 0.38 else (11 if alt >= 0.33 else 10))
    gf = s.shapes.add_table(n_lin, len(cab), Inches(0.35), Inches(1.15), Inches(sum(larg)), Inches(alt * n_lin))
    t = gf.table
    for i, w in enumerate(larg):
        t.columns[i].width = Inches(w)
    for r in t.rows:
        r.height = Inches(alt)
    # linha 0: grupos (a 1a coluna fica em branco)
    _celula(t.cell(0, 0), '', tam, 'FFFFFF', AZUL_ESCURO, True, PP_ALIGN.LEFT)
    c0 = 1
    for txt, n in grupos:
        t.cell(0, c0).merge(t.cell(0, c0 + n - 1))
        _celula(t.cell(0, c0), txt, tam, 'FFFFFF', AZUL_ESCURO, True, PP_ALIGN.CENTER)
        c0 += n
    for j, h in enumerate(cab):
        _celula(t.cell(1, j), h, tam, 'FFFFFF', AZUL_MEDIO, True, PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT)
    for i, lin in enumerate(linhas):
        fundo = 'FFFFFF' if i % 2 else CINZA_LINHA
        for j, v in enumerate(lin):
            _celula(t.cell(2 + i, j), v, tam, CINZA_TEXTO, fundo, False, PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT)
    if total:
        for j, v in enumerate(total):
            txt = v[0] if isinstance(v, tuple) else v
            _celula(t.cell(n_lin - 1, j), txt, tam, 'FFFFFF', AZUL_ESCURO, True, PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT)


def _resumo(prs, p):
    s = _base_conteudo(prs, p['secao'], p['titulo'], p.get('nota'))
    n = len(p['cards'])
    w, h, gap, x0, y0 = 2.0, 1.3, 0.12, 0.35, 1.15
    for i, (rot, val, nota, tom) in enumerate(p['cards']):
        x = x0 + i * (w + gap)
        _retangulo(s, x, y0, w, h, 'FFFFFF', MSO_SHAPE.ROUNDED_RECTANGLE, borda='D6DEE8').adjustments[0] = 0.08
        _texto(s, x + 0.12, y0 + 0.1, w - 0.24, 0.34, rot, 8, '6B7280', negrito=True)
        _texto(s, x + 0.12, y0 + 0.5, w - 0.24, 0.35, val, 17, TONS_CARD[tom], negrito=True)
        if nota:
            _texto(s, x + 0.12, y0 + 0.92, w - 0.24, 0.3, nota, 8, '6B7280')
    for i, (rot, val) in enumerate(p.get('extra') or []):
        x = x0 + i * 3.17
        _retangulo(s, x, 2.62, 3.05, 0.5, 'EAF2FB', MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.2
        _texto(s, x + 0.15, 2.66, 1.6, 0.42, rot, 9, '6B7280', negrito=True, ancora=MSO_ANCHOR.MIDDLE)
        _texto(s, x + 1.6, 2.66, 1.3, 0.42, val, 14, AZUL_ESCURO, negrito=True, alinhar=PP_ALIGN.RIGHT, ancora=MSO_ANCHOR.MIDDLE)
    _grafico_area(s, p, 0.35, 3.25, 12.6, 3.65)


def _grafico_area(s, p, x, y, w, h):
    cd = CategoryChartData()
    cd.categories = p['cats']
    for nome, vals, _ in p['series']:
        cd.add_series(nome, vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size, ch.legend.font.name = Pt(11), 'Calibri'
    ch.font.name, ch.font.size = 'Calibri', Pt(10)
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = _rgb('E3E8EF')
    va.format.line.fill.background()
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = '#,##0', False
    va.tick_labels.font.size = Pt(10)
    ch.category_axis.tick_labels.font.size = Pt(11)
    ch.plots[0].gap_width, ch.plots[0].overlap = 60, -5
    for ser, (_, _, cor) in zip(ch.plots[0].series, p['series']):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = _rgb(cor)
    real = ch.plots[0].series[-1]                                 # rotulos so na serie do ano atual
    real.data_labels.show_value = True
    real.data_labels.number_format, real.data_labels.number_format_is_linked = '#,##0', False
    real.data_labels.font.size, real.data_labels.font.bold = Pt(9), True
    real.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END


def _cartoes(s, cards, y0, h=1.2):
    n = len(cards)
    gap = 0.15
    w = (12.6 - gap * (n - 1)) / n
    for i, (rot, val, nota, tom) in enumerate(cards):
        x = 0.35 + i * (w + gap)
        _retangulo(s, x, y0, w, h, 'FFFFFF', MSO_SHAPE.ROUNDED_RECTANGLE, borda='D6DEE8').adjustments[0] = 0.08
        _texto(s, x + 0.14, y0 + 0.1, w - 0.28, 0.32, rot, 9, '6B7280', negrito=True)
        _texto(s, x + 0.14, y0 + 0.47, w - 0.28, 0.4, val, 20, TONS_CARD[tom], negrito=True)
        if nota:
            _texto(s, x + 0.14, y0 + 0.9, w - 0.28, 0.25, nota, 9, '6B7280')


def _grafico_pag(prs, p):
    s = _base_conteudo(prs, p['secao'], p['titulo'], p.get('nota'))
    topo = 1.15
    if p.get('cards'):
        _cartoes(s, p['cards'], topo)
        topo += 1.35
    for i, (rot, val) in enumerate(p.get('extra') or []):
        x = 0.35 + i * 3.17
        _retangulo(s, x, topo, 3.05, 0.45, 'EAF2FB', MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.2
        _texto(s, x + 0.12, topo + 0.03, 1.3, 0.39, rot, 8, '6B7280', negrito=True, ancora=MSO_ANCHOR.MIDDLE)
        _texto(s, x + 1.3, topo + 0.03, 1.65, 0.39, val, 13, AZUL_ESCURO, negrito=True, alinhar=PP_ALIGN.RIGHT, ancora=MSO_ANCHOR.MIDDLE)
    topo += 0.6
    _grafico_area(s, p, 0.35, topo, 12.6, 6.9 - topo)


def _cor_pct(v):
    return '1F9D55' if v >= 1 else ('D9922E' if v >= 0.8 else 'E04A4F')


def _intl(prs, p):
    """Tabela por pais (esquerda) e, junto, as barras de realizado x plano do mes e do acumulado (direita)."""
    s = _base_conteudo(prs, p['secao'], p['titulo'], p.get('nota'))
    tab = dict(p)
    _tabela(prs, tab, s)
    x0, w, h = 7.45, 5.5, 2.75
    for k, (chip, itens) in enumerate(p['barras']):
        y = 1.15 + k * (h + 0.15)
        _retangulo(s, x0, y, w, h, 'FFFFFF', MSO_SHAPE.ROUNDED_RECTANGLE, borda='D6DEE8').adjustments[0] = 0.04
        _texto(s, x0 + 0.2, y + 0.14, 3.4, 0.25, 'Receita mercado externo — por país', 11, AZUL_ESCURO, negrito=True)
        _retangulo(s, x0 + w - 1.65, y + 0.1, 1.5, 0.28, 'DCEBFA' if k == 0 else AZUL_ESCURO, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.5
        _texto(s, x0 + w - 1.65, y + 0.1, 1.5, 0.28, chip, 9, AZUL_ESCURO if k == 0 else 'FFFFFF', negrito=True, alinhar=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE)
        n = max(len(itens), 1)
        lh = min(0.3, (h - 0.6) / n)
        for i, (pais, real, plano) in enumerate(itens):
            yy = y + 0.55 + i * lh
            _texto(s, x0 + 0.2, yy, 1.1, lh, pais, 9, AZUL_ESCURO, negrito=True, ancora=MSO_ANCHOR.MIDDLE)
            tx, tw = x0 + 1.35, 2.2
            _retangulo(s, tx, yy + lh / 2 - 0.05, tw, 0.1, 'E8ECF2', MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.5
            frac = (real / plano) if plano else 0
            if plano:
                f = max(min(frac, 1.2) / 1.2, 0.012)
                _retangulo(s, tx, yy + lh / 2 - 0.05, tw * f, 0.1, _cor_pct(frac), MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.5
                _retangulo(s, tx + tw / 1.2, yy + lh / 2 - 0.1, 0.02, 0.2, '10263F')
            else:
                _retangulo(s, tx, yy + lh / 2 - 0.05, 0.03, 0.1, 'C3CAD6')
            _texto(s, x0 + 3.65, yy, 1.1, lh, f'R$ {real / 1000:,.0f} mil'.replace(',', '.'), 9, AZUL_ESCURO, negrito=True, alinhar=PP_ALIGN.RIGHT, ancora=MSO_ANCHOR.MIDDLE)
            _texto(s, x0 + 4.5, yy, 0.8, lh, (f'{frac * 100:.0f}%' if plano else '–'), 10, _cor_pct(frac) if plano else '6B7280', negrito=True, alinhar=PP_ALIGN.RIGHT, ancora=MSO_ANCHOR.MIDDLE)
        _texto(s, x0 + 0.2, y + h - 0.22, w - 0.4, 0.15, '● ≥100% do plano   ● 80–99%   ● <80%   ┃ = plano', 7, '6B7280')


def gerar(paginas: list[dict], gerado_em: str = '') -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    CTX.update(gerado=gerado_em, pag=0, total=len(paginas))
    for p in paginas:
        {'capa': _capa, 'resumo': _resumo, 'tabela': _tabela, 'intl': _intl, 'grafico': _grafico_pag}[p['tipo']](prs, p)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
