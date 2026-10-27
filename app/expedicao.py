"""Expedicao (Planejamento): NFs e estoque de expedicao, tudo lido do TOTVS (app/totvs/expedicao.py)."""
from flask import Blueprint, abort, jsonify, request
from flask_login import current_user, login_required

from app import db
from app.models import Category, Sector, Tab

exp_bp = Blueprint('expedicao', __name__, url_prefix='/expedicao')
SETORES = ('Diretoria', 'PCP', 'Produção', 'Comercial')


def garantir_aba() -> None:
    cat = Category.query.filter_by(slug='planejamento').first()
    if not cat or Tab.query.filter_by(slug='expedicao').first():
        return
    ordem = max([t.order or 0 for t in cat.tabs] or [0]) + 1
    tab = Tab(category_id=cat.id, slug='expedicao', name_pt='Expedição', name_es='Expedición', icon='\U0001F69A', icon_color='#0d9488', is_new=True, order=ordem)
    db.session.add(tab)
    tab.sectors = [s for s in Sector.query.all() if s.name in SETORES]
    db.session.commit()


def _autorizar():
    tab = Tab.query.filter_by(slug='expedicao').first_or_404()
    if not tab.visible_to(current_user):
        abort(403)


@exp_bp.route('/dados')
@login_required
def dados():
    _autorizar()
    from app.totvs import expedicao
    try:
        dias = max(1, min(int(request.args.get('dias', 30)), 90))
    except ValueError:
        dias = 30
    if not expedicao.ler().get('notas_em') and not expedicao.ESTADO['rodando']:     # ainda nao leu nada (servidor acabou de subir): le agora, sem esperar o agendador
        import threading
        threading.Thread(target=expedicao.atualizar, daemon=True).start()
    return jsonify(expedicao.painel(dias))


@exp_bp.route('/pedidos')
@login_required
def pedidos():
    _autorizar()
    from app.totvs import expedicao
    return jsonify(expedicao.pedidos_x_expedicao())


@exp_bp.route('/nf/buscar')
@login_required
def nf_buscar():
    _autorizar()
    from app.totvs import expedicao
    return jsonify(ok=True, notas=expedicao.buscar_nfs(request.args.get('q', '')))


@exp_bp.route('/nf/detalhe')
@login_required
def nf_detalhe():
    _autorizar()
    from app.totvs import expedicao
    try:
        return jsonify(ok=True, nota=expedicao.detalhe_nf(request.args.get('seq', ''), request.args.get('data', '')))
    except Exception as e:  # noqa: BLE001
        return jsonify(ok=False, erro=str(e)[:300])


@exp_bp.route('/nf/pdf')
@login_required
def nf_pdf():
    _autorizar()
    from flask import Response
    from app.totvs import expedicao
    try:
        n = expedicao.detalhe_nf(request.args.get('seq', ''), request.args.get('data', ''))
    except Exception as e:  # noqa: BLE001
        return Response(f'Não foi possível consultar a NF: {e}', status=502)
    nome = f"NF_{n['nf']}.pdf"
    try:
        return Response(expedicao.nf_pdf(n), mimetype='application/pdf', headers={'Content-Disposition': f'attachment; filename={nome}'})
    except Exception:  # noqa: BLE001
        import traceback
        traceback.print_exc()
        return Response(expedicao.nf_html(n).replace('</body>', '<script>window.onload=function(){window.print()}</script></body>'), mimetype='text/html')      # sem Chromium: abre a pagina pronta para "Salvar como PDF"


def _filtro_prontos(x, tipo):
    d = x['dias_exp'] or 0
    if tipo == 'prontos90' and d <= 90:
        return False
    fx = request.args.get('faixa', type=int)
    if fx is not None and not ((0, 30), (31, 60), (61, 90), (91, 10 ** 6))[max(0, min(fx, 3))][0] <= d <= ((0, 30), (31, 60), (61, 90), (91, 10 ** 6))[max(0, min(fx, 3))][1]:
        return False
    cl = request.args.get('cliente')
    return not cl or (x['cliente'] or '') == cl


@exp_bp.route('/exportar')
@login_required
def exportar():
    """Excel (.xlsx) das listas do painel: parados (+90 dias), faixa, inspecao, nfs (pendencias)."""
    _autorizar()
    import io
    from datetime import date
    from flask import send_file
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from app.totvs import expedicao
    tipo = request.args.get('tipo', 'parados')
    p = expedicao.painel(request.args.get('dias', 30, type=int)) if tipo not in ('prontos', 'prontos90', 'sem_pedido', 'hist') else {}
    E = p.get('estoque')
    dia = lambda v: date.fromisoformat(v).strftime('%d/%m/%Y') if v else ''
    if tipo == 'inspecao':
        I = E['inspecao']
        nome, cab = 'Aguardando inspeção', ['Produto', 'Descrição', f"Saldo em inspeção ({I['unidade']})", 'Saldo na expedição (m)', 'Peças na expedição']
        linhas = [[x['cod'], x['nome'], x['saldo'], x['exp_m'], x['exp_pecas']] for x in I['lista']]
    elif tipo in ('prontos', 'prontos90', 'sem_pedido', 'hist'):
        P = expedicao.pedidos_x_expedicao()
        if tipo in ('prontos', 'prontos90'):
            nome, cab = ('Prontos aguardando faturar' if tipo == 'prontos' else 'Prontos há mais de 90 dias'), ['Cliente', 'Unidade', 'Pedido', 'Artigo', 'Descrição', 'Data do pedido', 'Dias de pedido', 'Pendente a faturar (m)', 'Pronto na expedição (m)', 'Preço unitário (R$)', 'Valor pronto (R$)', 'Última entrada na expedição', 'Dias na expedição', 'Pedido → expedição (dias)']
            linhas = [[x['cliente'], x['unidade'], x['pedido'], x['cod'], x['nome'], dia(x['dt_pedido']), x['idade_pedido'], x['pendente'], x['pronto'], x['vl_unit'], x['valor'], dia(x['ult_entrada']) or 'sem entrada em 180 dias', x['dias_exp'], x['lead']] for x in P['prontos'] if _filtro_prontos(x, tipo)]
        elif tipo == 'hist':
            f = P['hist'][max(0, min(request.args.get('i', 0, type=int), len(P['hist']) - 1))]
            nome, cab = 'Pedido até expedição ' + f['faixa'], ['Artigo', 'Descrição', 'Metros', 'Peças', 'Pedido há (dias)', 'Prazo pedido → expedição (dias)', 'Na expedição há (dias)', 'Valor pronto (R$)']
            linhas = [[x['cod'], x['nome'], x['saldo'], x['pecas'], x['idade_pedido'], x['lead'], x['dias_exp'], x['valor']] for x in f['lista']]
        else:
            nome, cab = 'Estoque sem pedido', ['Artigo', 'Descrição', 'Metros', 'Peças', 'Dias desde a última entrada', 'Última entrada']
            linhas = [[x['cod'], x['nome'], x['saldo'], x['pecas'], x['dias_exp'], dia(x['ult_entrada']) or 'sem entrada em 180 dias'] for x in P['sem_pedido']]
    elif tipo == 'nfs':
        nome, cab = 'NFs com pendência', ['Data', 'Hora', 'NF', 'Série', 'Situação', 'Cliente', 'Metros', 'Valor (R$)']
        linhas = [[dia(n['dia']), n['hora'], n['nf'], n['serie'], n['situacao'], n['cliente'], n['metros'], n['valor']] for n in p['notas']['problemas']]
    else:
        if tipo == 'faixa':
            f = E['faixas'][max(0, min(request.args.get('i', 0, type=int), len(E['faixas']) - 1))]
            nome, lista = f['faixa'], f['lista']
        elif tipo == 'estoque':
            nome, lista = 'Estoque Expedição', [a for f in E['faixas'] for a in f['lista']]
        else:
            nome, lista = 'Parado há mais de 90 dias', E['parados90']
        cab = ['Artigo', 'Descrição', 'Metros', 'Peças', 'Dias sem movimento', 'Última movimentação', 'Última entrada (180 d)']
        linhas = [[a['cod'], a['nome'], a['saldo'], a['lotes'], ('mais de 180' if a['dias'] and a['dias'] > 180 else a['dias']), dia(a['ult_mov']), dia(a['ult_entrada']) or 'sem entrada em 180 dias'] for a in lista]
    wb = Workbook()
    ws = wb.active
    ws.title = nome[:31]
    ws.append(cab)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='0F2A4A')
    for l in linhas:
        ws.append(l)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(60, max(12, max(len(str(c.value or '')) for c in col) + 2))
    ws.freeze_panes = 'A2'
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f'expedicao_{tipo}_{date.today():%Y-%m-%d}.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@exp_bp.route('/diagnostico')
@login_required
def diagnostico():
    """So administrador: exemplo cru de lote e de NF para conferir os nomes dos campos do TOTVS."""
    if not current_user.is_admin:
        abort(403)
    from app.totvs import expedicao
    d = expedicao.ler()
    return jsonify(amostra_lote=d.get('amostra_lote'), amostra_nf=d.get('amostra_nf'), n_produtos=len(d.get('estoque') or {}), mov=dict(atualizado=expedicao.ler_mov().get('atualizado'), n=len(expedicao.ler_mov().get('produtos') or {}), erros=expedicao.ler_mov().get('erros'), amostra_erro=expedicao.ler_mov().get('amostra_erro')), n_notas=len(d.get('notas') or {}),
                   inspecao=dict(cod=d.get('inspecao_cod'), nome=d.get('inspecao_desc'), n=len(d.get('inspecao') or {}), total_itens=d.get('inspecao_total_itens'), saldos_vistos=d.get('estoques_vistos')),
                   flags_amostra=dict(list((d.get('flags') or {}).items())[:5]), status_nf=sorted({n[3] for n in (d.get('notas') or {}).values()}), status_nfe=sorted({n[4] for n in (d.get('notas') or {}).values()}),
                   estado=expedicao.ESTADO, amostra_item_pedido=__import__('app.totvs.pedidos', fromlist=['ESTADO']).ESTADO.get('amostra_item'), amostra_pedido=__import__('app.totvs.pedidos', fromlist=['ESTADO']).ESTADO.get('amostra_pedido'))
