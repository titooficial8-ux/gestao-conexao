"""Financeiro: titulos a receber/pagar, vencidos, exportacao, prorrogacao de vencimento e verificacao dos recursos do TOTVS (so leitura)."""
import io
import threading
from datetime import date

from flask import Blueprint, abort, jsonify, request, send_file
from flask_login import current_user, login_required

from app.models import Tab

fin_bp = Blueprint('financeiro', __name__, url_prefix='/financeiro')


def _autorizar():
    tab = Tab.query.filter_by(slug='painel-financeiro').first_or_404()
    if not tab.visible_to(current_user):
        abort(403)


@fin_bp.route('/dados')
@login_required
def dados():
    _autorizar()
    from app.totvs import financeiro
    if not financeiro.ler().get('atualizado') and not financeiro.ESTADO['rodando']:         # primeira vez: le agora
        threading.Thread(target=financeiro.atualizar, daemon=True).start()
    return jsonify(financeiro.painel())


@fin_bp.route('/verificacao')
@login_required
def verificacao():
    """O que o TOTVS oferece (boleto, alteracao de vencimento): so lista endpoints do swagger, nao chama nenhum."""
    _autorizar()
    from app.totvs import api, financeiro
    try:
        cli = api.TotvsModaClient()
        cli.autenticar()
        return jsonify(ok=True, **financeiro.verificacao(cli))
    except Exception as e:  # noqa: BLE001
        return jsonify(ok=False, erro=str(e)[:300])


@fin_bp.route('/boleto-teste', methods=['POST'])
@login_required
def boleto_teste():
    """So administrador: pede ao TOTVS o boleto de UMA parcela, para verificar se ele devolve o PDF."""
    if not current_user.is_admin:
        abort(403)
    from flask import Response
    from app.totvs import financeiro
    j = request.get_json(silent=True) or {}
    try:
        pdf, log = financeiro.boleto_totvs(str(j.get('doc', '')), str(j.get('parcela', '')), str(j.get('cod', '')), str(j.get('cnpj', '')))
    except Exception as e:  # noqa: BLE001
        return jsonify(ok=False, log=[str(e)[:300]])
    if pdf:
        return Response(pdf, mimetype='application/pdf', headers={'Content-Disposition': f"attachment; filename=boleto_{j.get('doc')}_{j.get('parcela')}.pdf"})
    return jsonify(ok=False, log=log)


@fin_bp.route('/diagnostico')
@login_required
def diagnostico():
    if not current_user.is_admin:
        abort(403)
    from app.totvs import financeiro
    d = financeiro.ler()
    return jsonify(url_receber=d.get('url_R'), url_pagar=d.get('url_P'), amostra_receber=d.get('amostra_R'), amostra_pagar=d.get('amostra_P'), n_receber=len(d.get('receber', [])), n_pagar=len(d.get('pagar', [])), nomes_erro=d.get('nomes_erro'), nomes_resolvidos=len(d.get('nomes') or {}),
                   status_vistos=sorted({t['status'] for t in d.get('receber', [])})[:30], estado=financeiro.ESTADO)


@fin_bp.route('/exportar')
@login_required
def exportar():
    _autorizar()
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from app.totvs import financeiro
    p = financeiro.painel()
    tipo = request.args.get('tipo', 'receber')
    mes = request.args.get('mes', '')                      # AAAA-MM; vazio = tudo
    dia = lambda v: date.fromisoformat(v).strftime('%d/%m/%Y') if v else ''
    if tipo == 'prorrogacao':
        nome, cab = 'Prorrogação de vencimento', ['Cliente', 'Regra de pagamento', 'Fatura / Duplicata', 'Parcela', 'Vencimento atual', 'Novo vencimento', 'Dias a prorrogar', 'Valor (R$)', 'Já vencido?', 'Banco', 'Nosso número']
        linhas = [[x['cliente'], x['regra'], x['doc'], x['parcela'], dia(x['venc']), dia(x['novo']), x['dias'], x['valor'], 'Sim' if x['vencido'] else 'Não', x['banco'], x['nosso']] for x in p['prorrogacao'] if not mes or x['venc'][:7] == mes]
    elif tipo == 'alteracoes':
        nome, cab = 'Alterações nos títulos', ['Quando', 'Tipo', 'Cliente', 'Fatura / Duplicata', 'Parcela', 'De', 'Para', 'Valor (R$)']
        linhas = [[x['em'].replace('T', ' '), x['tipo'], x['nome'], x['doc'], x['parcela'], x['de'], x['para'], x['valor']] for x in p['alteracoes'] if not mes or x['em'][:7] == mes]
    else:
        lista = {'receber': p['vencidos_receber'], 'pagar': p['vencidos_pagar'], 'exportacao': p['exportacao']}.get(tipo, p['vencidos_receber'])
        if mes and tipo == 'exportacao':
            lista = [x for x in lista if (x['venc'] or '')[:7] == mes]
        nome = {'receber': 'Vencidos a receber', 'pagar': 'Vencidos a pagar', 'exportacao': 'A receber exportação'}.get(tipo, 'Vencidos')
        cab = ['Cliente / Fornecedor', 'Fatura / Duplicata', 'Parcela', 'Emissão', 'Vencimento', 'Dias de atraso', 'Valor original', 'Saldo em aberto', 'Moeda', 'Banco', 'Situação']
        linhas = [[x['nome'], x['doc'], x['parcela'], dia(x['emissao']), dia(x['venc']), x['atraso'], x['valor'], x['saldo'], x['moeda'], x['banco'], x['status']] for x in lista]
    wb = Workbook()
    ws = wb.active
    ws.title = nome[:31]
    ws.append(cab)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='0F2A4A')
    for l in linhas:
        ws.append(l)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(50, max(12, max(len(str(c.value or '')) for c in col) + 2))
    ws.freeze_panes = 'A2'
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f'financeiro_{tipo}_{date.today():%Y-%m-%d}.xlsx', mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
