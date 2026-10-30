"""Estudo da Malharia: apresentacao (.pptx) no modelo da Conexao com a media de entrega, as maquinas que entregam/produzem bem (conforme o DS.METROS) e as trocas previstas.

Tudo sai da programacao atual do Gestao (Sequencia) e da conta de previsao (app/previsao.py): DS.METROS -> voltas (metragem x 100 / raport) -> RPM instalado x eficiencia -> pecas por 24 h.
So leitura: nao altera nada na programacao."""
from __future__ import annotations

import io
import os
import re
from collections import defaultdict
from datetime import date, datetime, timedelta

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from app.relatorio_fat import pptx_gen as G

ATIVAS = ('PRODUZINDO', 'EM_SEQUENCIA', 'PCP')
DIAS_TROCAS = 21
DIAS_HISTORICO = int(os.getenv('ESTUDO_DIAS_HISTORICO', '15'))        # quantos dias de producao (dias ativos) de cada maquina entram no historico
LINHAS_POR_SLIDE = 13


def _num(v, casas=1):
    return f'{v:,.{casas}f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def _cod(artigo) -> str:
    return (str(artigo or '').split(' ')[0]).strip()


def historico(agora: datetime, dias_ativos: int = None) -> dict:
    """Entrega REAL de cada maquina nas suas ULTIMAS producoes: pecas pesadas no local 23 (Relatorio 32) nos ultimos `dias_ativos` dias em que a maquina produziu (o dia de producao vai de 06:20 a 06:19;
    o dia de hoje, ainda aberto, fica de fora). Peca e metro por dia = total desses dias / quantidade de dias, ou seja, quanto a maquina entrega quando esta produzindo.
    Tamanho da peca = metro padrao da pesagem (DS.METROS)."""
    from app.models import Report32Movement
    k = dias_ativos or DIAS_HISTORICO
    hoje_prod = ((agora - timedelta(hours=6, minutes=20)).date())
    desde = agora - timedelta(days=120)                                       # olha ate 120 dias para tras para achar os ultimos K dias ativos de cada maquina
    linhas = Report32Movement.query.filter(Report32Movement.dt_real >= desde + timedelta(hours=3), Report32Movement.cod_maquina.isnot(None)).all()
    dia_maq = defaultdict(lambda: defaultdict(lambda: dict(pecas=0, metros=0.0, kg=0.0, ds=defaultdict(lambda: [0, 0.0]))))
    for r in linhas:
        try:
            n = int(re.sub(r'\D', '', r.cod_maquina or ''))
        except ValueError:
            continue
        dia = ((r.dt_real - timedelta(hours=3)) - timedelta(hours=6, minutes=20)).date()
        if dia >= hoje_prod:
            continue
        d = dia_maq[n][dia]
        d['pecas'] += 1
        d['metros'] += r.qt_movimento or 0
        d['kg'] += r.qt_pesos or 0
        if r.metro_padrao:
            g = d['ds'][round(r.metro_padrao)]
            g[0] += 1
            g[1] += r.qt_movimento or r.metro_padrao
    por_maquina, ds = [], defaultdict(lambda: dict(pecas=0, metros=0.0, maq_dias=0, maquinas=set()))
    inicio_min = None
    for n in sorted(dia_maq):
        dias = sorted(dia_maq[n], reverse=True)[:k]                          # as ultimas K producoes (dias ativos) desta maquina
        if not dias:
            continue
        pecas = sum(dia_maq[n][x]['pecas'] for x in dias)
        metros = sum(dia_maq[n][x]['metros'] for x in dias)
        por_maquina.append(dict(maquina=n, pecas=pecas, metros=metros, kg=sum(dia_maq[n][x]['kg'] for x in dias), dias_ativos=len(dias), de=min(dias), ate=max(dias),
                                pecas_dia=pecas / len(dias), metros_dia=metros / len(dias), peca_media=(metros / pecas) if pecas else None))
        inicio_min = min(inicio_min, min(dias)) if inicio_min else min(dias)
        for x in dias:
            for tam, (c, mt) in dia_maq[n][x]['ds'].items():
                g = ds[tam]
                g['pecas'] += c
                g['metros'] += mt
                g['maq_dias'] += 1
                g['maquinas'].add(n)
    por_tamanho = [dict(ds=t, pecas=v['pecas'], maquinas=len(v['maquinas']), pecas_dia_maquina=v['pecas'] / v['maq_dias'], metros_dia_maquina=v['metros'] / v['maq_dias']) for t, v in sorted(ds.items()) if v['pecas'] >= 5]
    return dict(dias_ativos=k, por_maquina=por_maquina, por_tamanho=por_tamanho, total_pecas=sum(m['pecas'] for m in por_maquina), desde=inicio_min)


def coletar(agora: datetime = None) -> dict:
    """Le a programacao e calcula, por maquina e por OP, o que o estudo precisa."""
    from app import previsao
    from app.main.routes import resolve_metros_from_article
    from app.models import Machine, Sequencia
    agora = agora or datetime.now()
    por_maq = defaultdict(list)
    ds_cache = {}
    for q in Sequencia.query.join(Machine).filter(Sequencia.status.in_(ATIVAS)).all():
        cod = _cod(q.artigo)
        if cod not in ds_cache:
            try:
                ds_cache[cod] = float(str(resolve_metros_from_article(q.artigo)).replace(',', '.'))
            except (TypeError, ValueError):
                ds_cache[cod] = None
        ds = ds_cache[cod]
        pd = q.pecas_dia or None
        fim = q.prev_termino_em or (datetime.combine(q.termino_malharia, datetime.min.time()) if q.termino_malharia else None)
        por_maq[q.machine.number].append(dict(
            maquina=q.machine.number, op=q.op, artigo=q.artigo, cod=cod, cliente=q.cliente or '', status=q.status, ordem=q.sequence_order if q.sequence_order is not None else 10 ** 6, ds=ds,
            total=q.total_pcs or 0, produzidas=q.pecas_produzidas or 0, restam=max((q.total_pcs or 0) - (q.pecas_produzidas or 0), 0), solicitada=q.qt_solicitada or 0,
            pecas_dia=pd, metros_dia=(pd * ds) if (pd and ds) else None, voltas=q.prev_voltas, rpm=q.prev_rpm, raport=q.prev_raport, fonte=q.prev_fonte, fim=fim))
    maquinas = []
    for n in sorted(por_maq):
        ops = sorted(por_maq[n], key=lambda o: (o['status'] != 'PRODUZINDO', o['ordem']))
        com = [o for o in ops if o['pecas_dia']]
        comm = [o for o in ops if o['metros_dia']]
        fins = [o['fim'] for o in ops if o['fim']]
        maquinas.append(dict(
            maquina=n, ops=ops, n_ops=len(ops), restam=sum(o['restam'] for o in ops), solicitada=sum(o['solicitada'] for o in ops),
            media_pecas=(sum(o['pecas_dia'] for o in com) / len(com)) if com else None, media_metros=(sum(o['metros_dia'] for o in comm) / len(comm)) if comm else None,
            fim=max(fins) if fins else None, sem_base=len(ops) - len(com)))
    # trocas: cada vez que a proxima OP da fila e de outro artigo, a troca acontece quando a anterior termina
    trocas = []
    limite = agora + timedelta(days=DIAS_TROCAS)
    for m in maquinas:
        for a, b in zip(m['ops'], m['ops'][1:]):
            if a['cod'] and b['cod'] and a['cod'] != b['cod'] and a['fim'] and agora <= a['fim'] <= limite:
                trocas.append(dict(quando=a['fim'], maquina=m['maquina'], sai=a, entra=b))
    trocas.sort(key=lambda t: t['quando'])
    # por tamanho de peca (DS.METROS)
    por_ds = defaultdict(list)
    for m in maquinas:
        for o in m['ops']:
            if o['ds'] and o['pecas_dia']:
                por_ds[o['ds']].append(o)
    grupos_ds = [dict(ds=ds, ops=len(v), maquinas=len({o['maquina'] for o in v}), pecas_dia=sum(o['pecas_dia'] for o in v) / len(v), metros_dia=sum(o['metros_dia'] for o in v) / len(v),
                      min_peca=sum(1440 / o['pecas_dia'] for o in v) / len(v)) for ds, v in sorted(por_ds.items())]
    sem_base = [o for m in maquinas for o in m['ops'] if not o['pecas_dia']]
    return dict(agora=agora, hist=historico(agora), maquinas=maquinas, trocas=trocas, grupos_ds=grupos_ds, sem_base=sem_base, eficiencia=previsao.EFICIENCIA, troca_horas=previsao.TROCA_HORAS, fonte=previsao.FONTE)


# ------------------------------------------------------------------ paginas
def _slide(prs, secao, titulo, nota=''):
    return G._base_conteudo(prs, secao, titulo, nota)


def _tabela(prs, secao, titulo, cab, larguras, linhas, nota='', alinhar_esq=(0,)):
    s = _slide(prs, secao, titulo, nota)
    n = len(linhas) + 1
    alt = min(0.36, 5.55 / n)
    gf = s.shapes.add_table(n, len(cab), Inches(0.35), Inches(1.15), Inches(sum(larguras)), Inches(alt * n))
    t = gf.table
    for i, w in enumerate(larguras):
        t.columns[i].width = Inches(w)
    for r in t.rows:
        r.height = Inches(alt)
    for j, h in enumerate(cab):
        G._celula(t.cell(0, j), h, 11, 'FFFFFF', G.AZUL_MEDIO, True, PP_ALIGN.LEFT if j in alinhar_esq else PP_ALIGN.RIGHT)
    for i, lin in enumerate(linhas):
        fundo = 'FFFFFF' if i % 2 else G.CINZA_LINHA
        for j, v in enumerate(lin):
            G._celula(t.cell(1 + i, j), v, 11, G.CINZA_TEXTO, fundo, False, PP_ALIGN.LEFT if j in alinhar_esq else PP_ALIGN.RIGHT)
    return s


def _barras(s, titulo, cats, vals, cor, x, y, w, h, fmt='#,##0.0'):
    G._retangulo(s, x, y, w, h, 'FFFFFF', MSO_SHAPE.ROUNDED_RECTANGLE, borda='D6DEE8').adjustments[0] = 0.03
    G._texto(s, x + 0.2, y + 0.12, w - 0.4, 0.3, titulo, 12, G.AZUL_ESCURO, negrito=True)
    if not cats:
        G._texto(s, x + 0.3, y + h / 2 - 0.2, w - 0.6, 0.4, 'Sem dados para calcular (faltam raport, RPM ou DS.METROS).', 12, '6B7280', alinhar=PP_ALIGN.CENTER)
        return
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series(titulo, vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x + 0.1), Inches(y + 0.5), Inches(w - 0.2), Inches(h - 0.6), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.name, ch.font.size = 'Calibri', Pt(10)
    ch.category_axis.reverse_order = True
    ch.category_axis.tick_labels.font.size = Pt(10)
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    pl = ch.plots[0]
    pl.gap_width = 45
    ser = pl.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = G._rgb(cor)
    pl.has_data_labels = True
    pl.data_labels.number_format, pl.data_labels.number_format_is_linked = fmt, False
    pl.data_labels.font.size, pl.data_labels.font.bold = Pt(10), True
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END


def _colunas(s, titulo, cats, vals, cor, x, y, w, h, fmt='#,##0.0'):
    G._retangulo(s, x, y, w, h, 'FFFFFF', MSO_SHAPE.ROUNDED_RECTANGLE, borda='D6DEE8').adjustments[0] = 0.03
    G._texto(s, x + 0.2, y + 0.12, w - 0.4, 0.3, titulo, 12, G.AZUL_ESCURO, negrito=True)
    if not cats:
        G._texto(s, x + 0.3, y + h / 2 - 0.2, w - 0.6, 0.4, 'Sem dados para calcular (faltam raport, RPM ou DS.METROS).', 12, '6B7280', alinhar=PP_ALIGN.CENTER)
        return
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series(titulo, vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x + 0.1), Inches(y + 0.5), Inches(w - 0.2), Inches(h - 0.6), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.name, ch.font.size = 'Calibri', Pt(10)
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    pl = ch.plots[0]
    pl.gap_width = 60
    pl.series[0].format.fill.solid()
    pl.series[0].format.fill.fore_color.rgb = G._rgb(cor)
    pl.has_data_labels = True
    pl.data_labels.number_format, pl.data_labels.number_format_is_linked = fmt, False
    pl.data_labels.font.size, pl.data_labels.font.bold = Pt(10), True
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END


def _bloco_texto(s, x, y, w, h, titulo, linhas, cor_titulo=None):
    G._retangulo(s, x, y, w, h, 'FFFFFF', MSO_SHAPE.ROUNDED_RECTANGLE, borda='D6DEE8').adjustments[0] = 0.03
    G._texto(s, x + 0.25, y + 0.15, w - 0.5, 0.3, titulo, 14, cor_titulo or G.AZUL_ESCURO, negrito=True)
    yy = y + 0.6
    for lin in linhas:
        G._texto(s, x + 0.25, yy, w - 0.5, 0.5, lin, 12, G.CINZA_TEXTO)
        yy += 0.55


def _planejar(d: dict) -> list:
    """Lista de paginas (funcoes que desenham), para o rodape saber o total."""
    maq = d['maquinas']
    com_dados = [m for m in maq if m['media_pecas']]
    paginas = []

    paginas.append(lambda prs: G._capa(prs, dict(titulo='Estudo da Malharia', sub='Entrega, produtividade e trocas')))

    def resumo(prs):
        s = _slide(prs, 'Estudo da Malharia', 'Resumo da programação', 'Conta: DS.METROS → voltas (raport) → RPM instalado × eficiência → peças por 24 h. Domingos e datas marcadas em Dias sem Produção não produzem.')
        todas = [o for m in maq for o in m['ops']]
        mp = [m['media_pecas'] for m in com_dados]
        mm = [m['media_metros'] for m in maq if m['media_metros']]
        fins = [m['fim'] for m in maq if m['fim']]
        G._cartoes(s, [('MÁQUINAS PROGRAMADAS', str(len(maq)), 'com OPs na fila', 'azul'), ('OPs NA PROGRAMAÇÃO', str(len(todas)), 'PCP, em sequência e produzindo', 'azul'),
                       ('PEÇAS A PRODUZIR', _num(sum(o['restam'] for o in todas), 0), 'o que ainda falta pesar', 'laranja'), ('QT. SOLICITADA', _num(sum(o['solicitada'] for o in todas), 0), 'soma das OPs', 'azul')], 1.15)
        G._cartoes(s, [('MÉDIA DE ENTREGA', _num(sum(mp) / len(mp), 1) if mp else '–', 'peças por 24 h, por máquina', 'verde'), ('MÉDIA DE METROS', _num(sum(mm) / len(mm), 0) if mm else '–', 'metros por 24 h, por máquina', 'verde'),
                       ('FIM DA FILA MAIS LONGA', max(fins).strftime('%d/%m/%Y') if fins else '–', 'última OP programada', 'laranja'), ('EFICIÊNCIA USADA', f'{d["eficiencia"] * 100:.0f}%', f'troca de artigo: {d["troca_horas"]:.0f} h', 'azul')], 2.65)
        G._texto(s, 0.35, 4.3, 12.6, 0.3, 'O que este estudo mostra', 15, G.AZUL_ESCURO, negrito=True)
        for i, t in enumerate(['Média de entrega: quantas peças cada máquina deve entregar por dia com os artigos da fila (previsão).',
                               'Máquinas que entregam bem e que produzem bem: histórico real de pesagens, por peças e por metros por dia (conforme o DS.METROS).',
                               'Trocas previstas: quando a máquina troca de artigo, nos próximos %d dias.' % DIAS_TROCAS,
                               'Pendências: OPs sem raport/RPM cadastrados, que ficam fora da média.']):
            G._texto(s, 0.55, 4.75 + i * 0.5, 12.2, 0.4, '•  ' + t, 13, G.CINZA_TEXTO)
    paginas.append(resumo)

    # media de entrega por maquina (tabela paginada)
    linhas_maq = [[f'MÁQ {m["maquina"]:02d}', str(m['n_ops']), _num(m['restam'], 0), _num(m['solicitada'], 0), _num(m['media_pecas'], 1) if m['media_pecas'] else '–',
                   _num(m['media_metros'], 0) if m['media_metros'] else '–', m['fim'].strftime('%d/%m/%Y') if m['fim'] else '–'] for m in maq]
    partes = [linhas_maq[i:i + LINHAS_POR_SLIDE + 1] for i in range(0, len(linhas_maq), LINHAS_POR_SLIDE + 1)] or [[]]
    for k, parte in enumerate(partes):
        paginas.append(lambda prs, parte=parte, k=k: _tabela(prs, 'Estudo da Malharia', 'Média de entrega por máquina' + (f' ({k + 1}/{len(partes)})' if len(partes) > 1 else ''),
                                                             ['Máquina', 'OPs', 'Peças a produzir', 'Qt. solicitada', 'Peças / 24 h (média)', 'Metros / 24 h (média)', 'Fim da fila'], [1.5, 1.0, 1.9, 1.9, 2.3, 2.3, 1.7], parte,
                                                             'Média simples das OPs da fila da máquina; OP sem raport/RPM fica de fora da média.'))

    H = d['hist']
    hm = [m for m in H['por_maquina'] if m['pecas_dia']]
    periodo = f'últimas {H["dias_ativos"]} produções (dias em que a máquina produziu) de cada máquina'

    def ranking(prs):
        s = _slide(prs, 'Estudo da Malharia', 'Máquinas que entregam e que produzem bem · histórico', f'Histórico real de pesagens (local 23, Relatório 32), {periodo}. Peças e metros por dia = o que a máquina entrega nos dias em que produz.')
        top_p = sorted(hm, key=lambda m: -m['pecas_dia'])[:10]
        top_m = sorted(hm, key=lambda m: -m['metros_dia'])[:10]
        _barras(s, 'Entregam bem · peças por dia', [f'MÁQ {m["maquina"]:02d}' for m in top_p], [round(m['pecas_dia'], 1) for m in top_p], '1B7F3B', 0.35, 1.15, 6.2, 5.75)
        _barras(s, 'Produzem bem · metros por dia', [f'MÁQ {m["maquina"]:02d}' for m in top_m], [round(m['metros_dia']) for m in top_m], '2E75B6', 6.75, 1.15, 6.2, 5.75, '#,##0')
    paginas.append(ranking)

    linhas_h = [[f'MÁQ {m["maquina"]:02d}', _num(m['pecas'], 0), _num(m['metros'], 0), str(m['dias_ativos']), _num(m['pecas_dia'], 1) if m['pecas_dia'] else '–', _num(m['metros_dia'], 0) if m['metros_dia'] else '–',
                 f'{_num(m["peca_media"], 0)} m' if m['peca_media'] else '–'] for m in sorted(hm, key=lambda m: -(m['metros_dia'] or 0))]
    partes_h = [linhas_h[i:i + LINHAS_POR_SLIDE + 1] for i in range(0, len(linhas_h), LINHAS_POR_SLIDE + 1)] or [[]]
    for k, parte in enumerate(partes_h):
        paginas.append(lambda prs, parte=parte, k=k: _tabela(prs, 'Estudo da Malharia', 'Histórico de entrega por máquina' + (f' ({k + 1}/{len(partes_h)})' if len(partes_h) > 1 else ''),
                                                             ['Máquina', 'Peças', 'Metros', 'Dias produzidos', 'Peças / dia', 'Metros / dia', 'Peça média'], [1.5, 1.4, 1.8, 2.2, 1.9, 1.9, 1.9], parte,
                                                             f'Histórico real, {periodo}. Ordenado por metros por dia. Peças e metros por dia = total ÷ dias em que a máquina produziu.'))

    def por_tamanho(prs):
        s = _slide(prs, 'Estudo da Malharia', 'Entrega por tamanho de peça (DS.METROS) · histórico', f'Histórico real, {periodo}. Peças e metros por dia = o que uma máquina entrega, em média, nos dias em que produziu esse tamanho.')
        g = H['por_tamanho']
        gf = s.shapes.add_table(len(g) + 1, 5, Inches(0.35), Inches(1.15), Inches(6.3), Inches(0.42 * (len(g) + 1)))
        t = gf.table
        for i, w in enumerate([1.3, 1.0, 1.1, 1.45, 1.45]):
            t.columns[i].width = Inches(w)
        for j, h in enumerate(['DS.METROS', 'Peças', 'Máquinas', 'Peças / dia', 'Metros / dia']):
            G._celula(t.cell(0, j), h, 11, 'FFFFFF', G.AZUL_MEDIO, True, PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT)
        for i, x in enumerate(g):
            for j, v in enumerate([f'{_num(x["ds"], 0)} m', _num(x['pecas'], 0), str(x['maquinas']), _num(x['pecas_dia_maquina'], 1), _num(x['metros_dia_maquina'], 0)]):
                G._celula(t.cell(1 + i, j), v, 12, G.CINZA_TEXTO, 'FFFFFF' if i % 2 else G.CINZA_LINHA, False, PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT)
        if not g:
            G._texto(s, 0.35, 1.2, 12.6, 0.4, 'Sem pesagens no período para calcular.', 13, '6B7280')
        else:
            _colunas(s, 'Peças por dia, por tamanho de peça', [f'{_num(x["ds"], 0)} m' for x in g], [round(x['pecas_dia_maquina'], 1) for x in g], '1B7F3B', 6.95, 1.15, 6.0, 2.75)
            _colunas(s, 'Metros por dia, por tamanho de peça', [f'{_num(x["ds"], 0)} m' for x in g], [round(x['metros_dia_maquina']) for x in g], '2E75B6', 6.95, 4.05, 6.0, 2.85, '#,##0')
    paginas.append(por_tamanho)

    # trocas
    tr = d['trocas']

    def trocas_resumo(prs):
        s = _slide(prs, 'Estudo da Malharia', f'Trocas previstas · próximos {DIAS_TROCAS} dias', 'Troca = a próxima OP da fila é de outro artigo; acontece quando a anterior termina (mais o tempo de troca da máquina).')
        agora = d['agora']
        em = lambda n: sum(1 for t in tr if t['quando'] <= agora + timedelta(days=n))
        por_maquina = defaultdict(int)
        for t in tr:
            por_maquina[t['maquina']] += 1
        G._cartoes(s, [('PRÓXIMOS 7 DIAS', str(em(7)), 'trocas de artigo', 'laranja'), ('PRÓXIMOS 14 DIAS', str(em(14)), 'trocas de artigo', 'laranja'), (f'PRÓXIMOS {DIAS_TROCAS} DIAS', str(len(tr)), 'trocas de artigo', 'azul'),
                       ('MÁQUINAS COM TROCA', str(len(por_maquina)), 'nesse período', 'azul')], 1.15)
        dias = [(agora + timedelta(days=i)).date() for i in range(DIAS_TROCAS)]
        cont = defaultdict(int)
        for t in tr:
            cont[t['quando'].date()] += 1
        _colunas(s, 'Trocas por dia', [x.strftime('%d/%m') for x in dias], [cont[x] for x in dias], 'B7791F', 0.35, 2.65, 12.6, 4.25, '0')
    paginas.append(trocas_resumo)
    linhas_tr = [[t['quando'].strftime('%d/%m %H:%M'), f'MÁQ {t["maquina"]:02d}', f'{t["sai"]["op"]} · {t["sai"]["artigo"]}'[:34], f'{t["entra"]["op"]} · {t["entra"]["artigo"]}'[:34],
                  (t['entra']['cliente'] or '–')[:16], f'{_num(t["entra"]["ds"], 0)} m' if t['entra']['ds'] else '–'] for t in tr]
    partes_tr = [linhas_tr[i:i + LINHAS_POR_SLIDE + 1] for i in range(0, len(linhas_tr), LINHAS_POR_SLIDE + 1)][:4]
    for k, parte in enumerate(partes_tr):
        paginas.append(lambda prs, parte=parte, k=k: _tabela(prs, 'Estudo da Malharia', 'Trocas de artigo previstas' + (f' ({k + 1}/{len(partes_tr)})' if len(partes_tr) > 1 else ''),
                                                             ['Quando', 'Máquina', 'Sai (OP · artigo)', 'Entra (OP · artigo)', 'Cliente', 'DS.METROS'], [1.6, 1.2, 3.5, 3.5, 1.7, 1.1], parte,
                                                             'Data e hora aproximadas, pelo término calculado da OP que sai.', alinhar_esq=(0, 1, 2, 3, 4)))

    sb = d['sem_base']
    if sb:
        linhas_sb = [[f'MÁQ {o["maquina"]:02d}', str(o['op']), (o['artigo'] or '')[:30], (o['cliente'] or '–')[:16], f'{_num(o["ds"], 0)} m' if o['ds'] else 'sem DS.METROS'] for o in sb[:LINHAS_POR_SLIDE + 1]]
        paginas.append(lambda prs: _tabela(prs, 'Estudo da Malharia', f'Pendências de cálculo · {len(sb)} OP(s) sem raport/RPM', ['Máquina', 'OP', 'Artigo', 'Cliente', 'DS.METROS'], [1.4, 1.4, 4.4, 3.2, 2.2], linhas_sb,
                                           'Cadastre o raport da estrutura e o RPM instalado da máquina (página Raport e RPM) para entrar na média.', alinhar_esq=(0, 1, 2, 3, 4)))

    def metodologia(prs):
        s = _slide(prs, 'Estudo da Malharia', 'Como o cálculo é feito', 'Parâmetros desta geração. Alterações de raport, RPM, DS.METROS ou dias sem produção mudam o estudo na próxima geração.')
        _bloco_texto(s, 0.35, 1.15, 6.2, 5.7, 'Informação técnica de produção', [
            '1. DS.METROS: tamanho da peça (METRO IDEAL do artigo, TOTVS).', '2. Voltas por peça = DS.METROS × 100 ÷ RAPORT (cm por volta, da estrutura).', '3. RPM efetivo = RPM instalado da máquina × eficiência.',
            '4. Minutos por peça = VOLTAS ÷ RPM efetivo.', '5. Peças por 24 h = 1.440 min ÷ minutos por peça.', '6. Metros por 24 h = peças por 24 h × DS.METROS.'])
        _bloco_texto(s, 6.75, 1.15, 6.2, 5.7, 'Parâmetros e regras', [
            f'Eficiência da máquina: {d["eficiencia"] * 100:.0f}% do RPM instalado.', f'Troca de artigo: {d["troca_horas"]:.0f} h antes de cada OP que ainda não está produzindo.', 'Domingo não produz; sábado produz.',
            'Datas marcadas em Dias sem Produção também não produzem.', 'A OP produzindo começa agora; cada próxima começa quando a anterior termina.',
            f'Fonte de voltas/RPM: {"raport da estrutura × RPM instalado" if d["fonte"] == "raport" else d["fonte"]}.'])
    paginas.append(metodologia)
    return paginas


def gerar(agora: datetime = None) -> bytes:
    d = coletar(agora)
    paginas = _planejar(d)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    G.CTX.update(gerado=d['agora'].strftime('%d/%m/%Y %H:%M'), pag=0, total=len(paginas), rotulo='Estudo da Malharia')
    try:
        for desenhar in paginas:
            desenhar(prs)
    finally:
        G.CTX['rotulo'] = 'Relatório de Faturamento'
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
