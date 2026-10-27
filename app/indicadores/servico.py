"""Dados da aba Indicadores.

Fontes (todas da API do TOTVS, so leitura):
- Malharia: as pecas pesadas no local 23 (baixas de fio, a mesma base do Controle de Producao).
- Acabamento: transacoes 2030/2031 (entrada da peca acabada), com metros e kg (metros x gramatura x largura / 1000).
- Faturamento: notas de venda (VDA) do Relatorio 150, metros = quantidade total do dia, ticket = valor bruto / metros.
Dia de producao = 06:20 ate 06:20 do dia seguinte; turnos 06:20-14:20, 14:20-22:20 e 22:20-06:20.
"""
from __future__ import annotations

import json
import os
import threading
import time
import traceback
from collections import defaultdict
from datetime import date, datetime, timedelta

from app import db
from app.indicadores.models import IndAcabamento, IndEstoqueItem, IndProduto, IndProdutoUn, IndReprovaCru
from app.totvs import api

BASE_URL = api.BASE_URL
TRX_URL = f'{BASE_URL}/api/totvsmoda/general/v2/transactions/search'
PROD_URL = f'{BASE_URL}/api/totvsmoda/product/v2/products/search'
EMPRESA = 1
OPS_ACABAMENTO = [int(x) for x in os.getenv('INDICADORES_OPS_ACABAMENTO', '2030,2031').split(',') if x.strip()]
DESDE = os.getenv('INDICADORES_DESDE', '2026-09-01')
INTERVALO_MIN = float(os.getenv('INDICADORES_INTERVALO_MIN', '2'))
INICIO_DIA = timedelta(hours=6, minutes=20)
ESTADO = dict(ultima=None, erro=None, pecas=0)
_trava = threading.Lock()
_cache: dict = {}
_vals: dict = {}                                                  # nome -> (quando, valor): resposta pronta, devolvida na hora
_calc_em: set = set()
_calc_trava = threading.Lock()


def _swr(nome, ttl, fn, *args):
    """Devolve na hora o ultimo valor calculado (mesmo velho) e recalcula em segundo plano quando passa de `ttl`; so a 1a vez espera.
    Se o calculo falha, a falha tambem fica guardada (e relancada na hora) ate a proxima tentativa em segundo plano: uma fonte fora do ar
    nao deixa a tela esperando o tempo limite a cada abertura."""
    from flask import current_app
    chave = (nome,) + args
    app = current_app._get_current_object()

    def calcular():
        try:
            with app.app_context():
                v = fn(*args)
            _vals[chave] = (time.time(), v, None)
        except Exception as e:  # noqa: BLE001
            if 'SEM_PLANILHA' not in str(e):
                traceback.print_exc()
            antigo = _vals.get(chave)
            _vals[chave] = (time.time(), antigo[1] if antigo else None, f'{type(e).__name__}: {str(e)[:160]}')
        finally:
            _calc_em.discard(chave)

    ent = _vals.get(chave)
    if ent is None:
        with _calc_trava:                                          # 1a vez: espera (uma so thread calcula)
            ent = _vals.get(chave)
            if ent is None:
                _calc_em.add(chave)
                calcular()
                ent = _vals.get(chave)
    elif time.time() - ent[0] > ttl and chave not in _calc_em:
        _calc_em.add(chave)
        threading.Thread(target=calcular, daemon=True, name='indicadores-calc').start()
    if ent[2] and ent[1] is None:
        raise RuntimeError(ent[2])
    return ent[1]


def dia_producao(h: datetime) -> date:
    return (h - INICIO_DIA).date()


def turno(h: datetime) -> int:
    x = h - INICIO_DIA
    return min((x.hour * 3600 + x.minute * 60 + x.second) // 28800, 2)


# ---------------------------------------------------------------- acabamento (TOTVS)
def _num(v):
    try:
        return float(str(v).replace(',', '.'))
    except (TypeError, ValueError):
        return None


def _buscar_trx(cli, ini: date, fim: date) -> list:
    saida = []
    a = ini
    while a <= fim:
        b = min(fim, a + timedelta(days=13))
        pag = 1
        while True:
            r = cli.post(TRX_URL, dict(branchCodeList=[EMPRESA], operationCodeList=OPS_ACABAMENTO, startChangeDate=f'{a.isoformat()}T00:00:00',
                                       endChangeDate=f'{(b + timedelta(days=1)).isoformat()}T00:00:00', page=pag, pageSize=500))
            if r.status_code != 200:
                raise RuntimeError(f'transacoes {a}..{b}: HTTP {r.status_code} {r.text[:160]}')
            j = r.json()
            saida += j.get('items') or []
            if not j.get('hasNext') or pag >= 80:
                break
            pag += 1
        a = b + timedelta(days=1)
    return saida


def _produtos(cli, codigos: list) -> None:
    """Gramatura e largura dos produtos que ainda nao conhecemos."""
    faltam = [c for c in codigos if not db.session.get(IndProduto, c)]
    for i in range(0, len(faltam), 50):
        lote = faltam[i:i + 50]
        r = cli.post(PROD_URL, dict(filter=dict(productCodeList=lote), option=dict(branchInfoCode=EMPRESA), page=1, pageSize=50, expand='additionalFields'))
        if r.status_code != 200:
            raise RuntimeError(f'produtos: HTTP {r.status_code} {r.text[:160]}')
        achados = {}
        for p in r.json().get('items') or []:
            campos = {str(f.get('name') or '').upper(): f.get('value') for f in p.get('additionalFields') or []}
            achados[p['productCode']] = IndProduto(codigo=p['productCode'], nome=(p.get('productName') or '')[:200],
                                                   gramatura=_num(campos.get('GRAMATURA M2')), largura=_num(campos.get('LARGURA PADRAO')))
        for c in lote:
            db.session.add(achados.get(c) or IndProduto(codigo=c, nome='', gramatura=None, largura=None))
        db.session.commit()


def atualizar_acabamento(app) -> bool:
    """Busca as pecas novas do acabamento (so leitura) e calcula o kg de cada uma."""
    if not _trava.acquire(blocking=False):
        return False
    try:
        with app.app_context():
            cli = api.TotvsModaClient()
            cli.autenticar()
            ultimo = db.session.query(db.func.max(IndAcabamento.mudou)).scalar()
            ini = (ultimo.date() - timedelta(days=2)) if ultimo else date.fromisoformat(DESDE)
            hoje = date.today()
            itens = _buscar_trx(cli, ini, hoje)
            existentes = {c for (c,) in db.session.query(IndAcabamento.codigo).all()}          # todos: uma transacao reprocessada no TOTVS volta com data nova mas e a mesma
            novos = []
            for t in itens:
                cod = t.get('transactionCode')
                if cod in existentes or not t.get('lastchangeDate'):
                    continue
                for it in t.get('items') or []:
                    q = _num(it.get('quantity'))
                    if not q or it.get('productCode') is None:
                        continue
                    novos.append(IndAcabamento(codigo=cod, operacao=t.get('operationCode'), mudou=datetime.fromisoformat(t['lastchangeDate'][:19]), produto=it['productCode'], metros=q))
                    existentes.add(cod)
                    break
            if novos:
                db.session.add_all(novos)
                db.session.commit()
            _produtos(cli, sorted({r.produto for r in IndAcabamento.query.filter(IndAcabamento.kg.is_(None)).all()}))
            fator = {p.codigo: (p.gramatura * p.largura / 1000 if p.gramatura and p.largura else None) for p in IndProduto.query.all()}
            for r in IndAcabamento.query.filter(IndAcabamento.kg.is_(None)).all():
                if fator.get(r.produto):
                    r.kg = round(r.metros * fator[r.produto], 3)
            db.session.commit()
            ESTADO.update(ultima=datetime.now().isoformat(timespec='seconds'), erro=None, pecas=IndAcabamento.query.count())
            _cache.clear()
        return True
    except Exception as e:  # noqa: BLE001
        traceback.print_exc()
        try:
            db.session.rollback()
        except Exception:  # noqa: BLE001
            pass
        ESTADO['erro'] = str(e)[:200]
        return False
    finally:
        _trava.release()


def iniciar_agendador(app) -> None:
    def loop():
        time.sleep(15)
        aquecer(app)                                                  # tela pronta logo na partida
        while True:
            try:
                atualizar_acabamento(app)
                atualizar_reprova_cru(app)
                atualizar_estoque(app)
                aquecer(app)
            except Exception:  # noqa: BLE001
                traceback.print_exc()
            time.sleep(INTERVALO_MIN * 60)
    threading.Thread(target=loop, daemon=True, name='indicadores-acabamento').start()


# ---------------------------------------------------------------- agregacao por dia de producao e turno
def _vazio():
    return {'t': [dict(pecas=0, metros=0.0, kg=0.0, m_real=0.0, m_padrao=0.0, sem_metros=0) for _ in range(3)]}


def _fechar(d: dict) -> dict:
    d['total'] = dict(pecas=sum(t['pecas'] for t in d['t']), metros=round(sum(t['metros'] for t in d['t']), 1), kg=round(sum(t['kg'] for t in d['t']), 1),
                      m_real=round(sum(t.get('m_real', 0) for t in d['t']), 1), m_padrao=round(sum(t.get('m_padrao', 0) for t in d['t']), 1), sem_metros=sum(t.get('sem_metros', 0) for t in d['t']))
    for t in d['t']:
        t['metros'] = round(t['metros'], 1)
        t['kg'] = round(t['kg'], 1)
    return d


def _malharia_calc() -> dict:
    """{data: {t:[3 turnos], total}} das pecas pesadas no local 23 (sem amostras), com a mesma conta do Controle de Producao."""
    from app.main.routes import _api_baixas, _hub_movimentos_json
    linhas = json.loads('[' + _api_baixas(dict(_hub_movimentos_json())) + ']')
    dias = defaultdict(_vazio)
    for r in linhas:
        if r.get('cod_artigo') and 'amostra' in r['cod_artigo'].lower():
            continue
        h = datetime.fromisoformat(r['dt_real'].replace('Z', '')) - timedelta(hours=3)
        alvo = dias[dia_producao(h)]['t'][turno(h)]
        alvo['pecas'] += 1
        real, padrao = r.get('qt_movimento') or 0, r.get('metro_padrao') or 0
        alvo['metros'] += padrao or real                          # como o Controle de Producao: metro padrao do artigo (o metro real fica na conferencia)
        alvo['m_real'] += real
        alvo['m_padrao'] += padrao
        alvo['sem_metros'] += 0 if (real or padrao) else 1
        alvo['kg'] += r.get('qt_pesos') or 0
    saida = {d.isoformat(): _fechar(v) for d, v in dias.items()}
    return saida


def _acabamento_calc() -> dict:
    dias = defaultdict(_vazio)
    for r in IndAcabamento.query.order_by(IndAcabamento.mudou).all():
        alvo = dias[dia_producao(r.mudou)]['t'][turno(r.mudou)]
        alvo['pecas'] += 1
        alvo['metros'] += r.metros or 0
        alvo['kg'] += r.kg or 0
        alvo['sem_metros'] += 0 if r.kg else 1                     # aqui 'sem_metros' = pecas sem kg (artigo sem gramatura/largura no cadastro)
    return {d.isoformat(): _fechar(v) for d, v in dias.items()}


def _faturamento_calc(ano: int, mes: int) -> dict:
    """Metros = quantidade total das vendas (VDA) do dia; ticket = valor bruto / metros (como a planilha FATURAMENTO)."""
    from app.totvs import base, servico
    cols, linhas, _ = servico.linhas_armazem()
    if not cols:
        return {}
    metros, valor = defaultdict(float), defaultdict(float)
    for l in linhas:
        x = base.derivar(dict(zip(cols, l)))
        if x['tipo'] != 'VDA' or x['ano'] != ano or x['mes'] != mes or not x['data']:
            continue
        d = x['data'].date().isoformat() if hasattr(x['data'], 'date') else str(x['data'])[:10]
        metros[d] += x['quant']
        valor[d] += x['bruta']
    try:                                                           # o Relatorio 150 e atualizado 1x por dia (08:00): o dia de hoje vem direto das NFs lidas a cada 2 min (Expedicao)
        from app.totvs import expedicao
        hoje = date.today()
        for d, (m, v) in expedicao.vendas_por_dia(ano, mes).items():
            if d == hoje.isoformat() or d not in metros:
                metros[d], valor[d] = m, v
    except Exception:  # noqa: BLE001
        traceback.print_exc()
    return {d: dict(metros=round(metros[d], 1), valor=round(valor[d], 2), ticket=round(valor[d] / metros[d], 3) if metros[d] else 0) for d in sorted(metros)}


def malharia_por_dia() -> dict:
    return _swr('malharia', 60, _malharia_calc)


def acabamento_por_dia() -> dict:
    return _swr('acabamento', 60, _acabamento_calc)


def faturamento_por_dia(ano: int, mes: int) -> dict:
    return _swr('faturamento', 120, _faturamento_calc, ano, mes)


def aquecer(app) -> None:
    """Calcula tudo antes de alguem pedir (chamado a cada ciclo): a tela abre e troca de dia sem esperar."""
    with app.app_context():
        for fn in (malharia_por_dia, acabamento_por_dia):
            try:
                fn()
            except Exception:  # noqa: BLE001
                traceback.print_exc()
        hoje = date.today()
        try:
            scrap_do_dia((datetime.now() - INICIO_DIA).date())
        except Exception:  # noqa: BLE001
            pass
        try:
            setups_do_dia((datetime.now() - INICIO_DIA).date())
        except Exception:  # noqa: BLE001
            pass
        try:
            faturamento_por_dia(hoje.year, hoje.month)
        except Exception:  # noqa: BLE001
            traceback.print_exc()


def _seguro(fonte, fn, avisos: list, *args):
    """Roda uma fonte de dados; se falhar, devolve vazio e guarda o motivo (a tela mostra o aviso em vez de dar erro 500)."""
    try:
        return fn(*args)
    except Exception as e:  # noqa: BLE001
        traceback.print_exc()
        avisos.append(f'{fonte}: {type(e).__name__}: {str(e)[:160]}')
        return {}


MEDIA_DESDE = os.getenv('INDICADORES_MEDIA_DESDE', DESDE)        # a "Media Bruta" da planilha e corrida desde o inicio da base (01/09/2026), nao so do mes


def _media(serie: dict, ate: str, minimo: float = 0.0) -> float:
    """Media Bruta como na planilha (aba BASE DIA): media dos metros por dia desde o inicio da base ate o dia escolhido, so dos dias COM producao
    (domingo/feriado sem producao nao entra). `minimo` tira do calculo o dia de partida incompleto."""
    v = [x['total']['metros'] for d, x in serie.items() if MEDIA_DESDE <= d <= ate and x['total']['metros'] > minimo]
    return round(sum(v) / len(v), 1) if v else 0.0


def reuniao(dia: date) -> dict:
    ch = dia.isoformat()
    avisos = []
    mal, aca = _seguro('Malharia', malharia_por_dia, avisos), _seguro('Acabamento', acabamento_por_dia, avisos)
    fat = _seguro('Faturamento', faturamento_por_dia, avisos, dia.year, dia.month)
    m, a = mal.get(ch) or _fechar(_vazio()), aca.get(ch) or _fechar(_vazio())
    f = fat.get(ch) or dict(metros=0, valor=0, ticket=0)
    tm = sum(x['metros'] for x in fat.values())
    tv = sum(x['valor'] for x in fat.values())
    return dict(
        data=ch, semana=(dia.timetuple().tm_yday - 1) // 7 + 1,
        malharia=dict(**m, media=_media(mal, ch), **_setups_resumo(_seguro('Setups', setups_do_dia, avisos, dia))),
        acabamento=dict(**a, media=_media(aca, ch, 2000)),
        faturamento=dict(metros=f['metros'], ticket_dia=f['ticket'], ticket_mes=round(tv / tm, 3) if tm else 0,
                         pct_prod=round(f['metros'] / m['total']['metros'] - 1, 4) if m['total']['metros'] else None,
                         pct_aca=round(f['metros'] / a['total']['metros'] - 1, 4) if a['total']['metros'] else None),
        estoque=_seguro('Estoque', estoque_dia, avisos, dia) or dict(entrada_total=dict(pecas=0, metros=0), saldo=dict(pecas=0, metros=0), saidas=dict(pecas=0, metros=0), entradas=[], monitorando_desde=None),
        scrap=_scrap_ou_none(avisos, dia),
        reprova_cru=_seguro('Reprova Cru', reprova_cru_dia, avisos, dia) or dict(pecas=0, metros=0, kg=0, linhas=[], turnos=[]),
        atualizado=ESTADO['ultima'], erro=ESTADO['erro'], avisos=avisos, intervalo_min=INTERVALO_MIN)


def _mes_anterior(ano: int, mes: int):
    return (ano - 1, 12) if mes == 1 else (ano, mes - 1)


def _serie_anterior(ano: int, mes: int, ate_dia: int, mal: dict, aca: dict, fat: dict) -> dict:
    """Metros por dia do mes anterior (ate o mesmo dia do mes) para comparar o acumulado."""
    def dias(origem, chave):
        out = []
        for d in range(1, ate_dia + 1):
            ch = date(ano, mes, d).isoformat() if d <= 28 or _dia_valido(ano, mes, d) else None
            v = origem.get(ch) if ch else None
            out.append(((v['total']['metros'] if chave != 'fat' else v['metros']) if v else 0.0))
        return out
    return dict(ano=ano, mes=mes, malharia=dias(mal, 'm'), acabamento=dias(aca, 'a'), faturamento=dias(fat, 'fat'))


def _dia_valido(ano: int, mes: int, d: int) -> bool:
    try:
        date(ano, mes, d)
        return True
    except ValueError:
        return False


def mensal(ano: int, mes: int) -> dict:
    ini = date(ano, mes, 1)
    fim = (date(ano + (mes == 12), mes % 12 + 1, 1))
    avisos = []
    mal, aca = _seguro('Malharia', malharia_por_dia, avisos), _seguro('Acabamento', acabamento_por_dia, avisos)
    fat = _seguro('Faturamento', faturamento_por_dia, avisos, ano, mes)
    dias = []
    d = ini
    while d < min(fim, date.today() + timedelta(days=1)):
        ch = d.isoformat()
        dias.append(dict(data=ch, malharia=mal.get(ch) or _fechar(_vazio()), acabamento=aca.get(ch) or _fechar(_vazio()), faturamento=fat.get(ch) or dict(metros=0, valor=0, ticket=0)))
        d += timedelta(days=1)
    pa, pm = _mes_anterior(ano, mes)
    fat_ant = _seguro('Faturamento (mês anterior)', faturamento_por_dia, avisos, pa, pm)
    anterior = _serie_anterior(pa, pm, len(dias) or 1, mal, aca, fat_ant)
    return dict(ano=ano, mes=mes, dias=dias, anterior=anterior, avisos=avisos)


# ---------------------------------------------------------------- reprova cru (NCs do SharePoint, local 24)
DEFEITOS = {
    'R01': 'AGULHA CILINDRO', 'R02': 'BARRAMENTO DE FIO', 'R03': 'BARRAMENTO DE TINTURARIA', 'R04': 'CONTAMINAÇÃO DE FIO', 'R05': 'AGULHA DISCO',
    'R06': 'FIO PUXADO', 'R07': 'FIO QUEBRADO', 'R08': 'MALHA CORRIDA CILINDRO', 'R09': 'MALHA CORRIDA DISCO', 'R10': 'MANCHA DE TINTURARIA',
    'R11': 'RISCO DE AGULHA', 'R12': 'FIO ERRADO', 'R13': 'RISCO DE ÓLEO', 'R14': 'TONALIDADE ERRADA', 'R16': 'SOBREPOSTO', 'R18': 'MANCHA DE ÓLEO',
    'R19': 'ESTOURO', 'R20': 'PLATINA', 'R21': 'FALHA DE DESENHO', 'R23': 'AGULHA PICANDO', 'R25': 'ATUADOR', 'R26': 'AMASSADO',
    'R27': 'PILLING (BOLINHAS)', 'R28': 'REBOBINAR', 'R29': 'SUJEIRA', 'R30': 'FORA DE LARGURA', 'R31': 'RISCO NO VINCO', 'R32': 'PARADA DE MÁQUINA',
    'R34': 'ENCHIMENTO SAINDO NO VINCO', 'R35': 'QUEBRADURA', 'REGULAGEM': 'REGULAGEM'}
_TURNO_NC = {'A': 1, 'B': 2, 'C': 3}
_ult_nc = {'ts': 0.0, 'erro': None, 'completo': 0.0}


def _linha_nc(x: dict):
    import re
    d = str(x.get('data1') or '')[:10]
    if not re.match(r'\d\d/\d\d/\d{4}$', d):
        return None
    dia = datetime.strptime(d, '%d/%m/%Y').date()
    metros = _num(x.get('Metros2'))
    if dia.year < 2026 or not metros:
        return None
    codigos = [c for c in re.split(r'\s+', str(x.get('contagemR') or '').strip().upper()) if c]
    return dict(sp_id=int(x['_id']), data=dia, maquina=str(x.get('Maquina') or '').strip().lstrip('0') or None, artigo=str(x.get('field_7') or '').strip()[:30],
                op=str(x.get('Ordemdeproducao') or '').strip()[:20], cliente=str(x.get('field_8') or '').strip()[:120], metros=metros, kg=_num(x.get('Quilos2')) or 0.0,
                turno=_TURNO_NC.get(str(x.get('TURNO') or '').strip().upper()), codigos=','.join(codigos)[:80],
                defeito=DEFEITOS.get(codigos[0], codigos[0]) if codigos else '', emitente=str(x.get('field_2') or '')[:80],
                criado=datetime.fromisoformat(str(x['_criado']).replace('Z', '')) if x.get('_criado') else None)


def atualizar_reprova_cru(app) -> bool:
    """Le as NCs do SharePoint (so leitura) e atualiza a tabela local; so entram NCs de 2026 em diante com metros e data."""
    from app.indicadores import sharepoint
    if not sharepoint.configurado():
        _ult_nc['erro'] = 'SharePoint não configurado no .env: ' + ', '.join(sharepoint.faltando())
        return False
    try:
        completo = time.time() - _ult_nc['completo'] > 3600 or not _ult_nc['ts']          # 1x por hora le tudo (para notar NCs apagadas); nas outras so o que mudou
        if completo:
            itens = sharepoint.itens(limite=30000)
        else:
            corte = datetime.utcfromtimestamp(_ult_nc['ts'] - 3600).strftime('%Y-%m-%dT%H:%M:%SZ')
            itens = sharepoint.itens_recentes(corte)
        linhas = [r for r in (_linha_nc(x) for x in itens) if r]
        with app.app_context():
            existentes = {r.sp_id: r for r in IndReprovaCru.query.all()}
            for r in linhas:
                alvo = existentes.get(r['sp_id'])
                if alvo is None:
                    db.session.add(IndReprovaCru(**r))
                else:
                    for k, v in r.items():
                        setattr(alvo, k, v)
            if completo:
                vivos = {r['sp_id'] for r in linhas}
                for sp_id, alvo in existentes.items():
                    if sp_id not in vivos:                          # NC apagada ou alterada para fora do criterio
                        db.session.delete(alvo)
            db.session.commit()
        _ult_nc.update(ts=time.time(), erro=None)
        if completo:
            _ult_nc['completo'] = time.time()
        _cache.clear()
        return True
    except Exception as e:  # noqa: BLE001
        traceback.print_exc()
        try:
            with app.app_context():
                db.session.rollback()
        except Exception:  # noqa: BLE001
            pass
        _ult_nc['erro'] = str(e)[:200]
        return False


def _nc_dict(r) -> dict:
    return dict(data=r.data.isoformat(), maquina=r.maquina, artigo=r.artigo, op=r.op, cliente=r.cliente, metros=round(r.metros, 1), kg=round(r.kg, 2),
                turno=r.turno, codigos=r.codigos, defeito=r.defeito)


def reprova_cru_dia(dia: date) -> dict:
    rs = IndReprovaCru.query.filter_by(data=dia).order_by(IndReprovaCru.criado).all()
    return dict(pecas=len(rs), metros=round(sum(r.metros for r in rs), 1), kg=round(sum(r.kg for r in rs), 1), linhas=[_nc_dict(r) for r in rs],
                turnos=[dict(pecas=sum(1 for r in rs if r.turno == t), metros=round(sum(r.metros for r in rs if r.turno == t), 1), kg=round(sum(r.kg for r in rs if r.turno == t), 1)) for t in (1, 2, 3)])


def reprova_cru_mensal(ano: int, mes: int) -> dict:
    ini = date(ano, mes, 1)
    fim = date(ano + (mes == 12), mes % 12 + 1, 1)
    rs = IndReprovaCru.query.filter(IndReprovaCru.data >= ini, IndReprovaCru.data < fim).order_by(IndReprovaCru.data, IndReprovaCru.criado).all()
    por_def = defaultdict(lambda: [0, 0.0, 0.0])
    por_dia = defaultdict(lambda: [0, 0.0, 0.0])
    for r in rs:
        for alvo in (por_def[r.defeito or 'SEM CÓDIGO'], por_dia[r.data.isoformat()]):
            alvo[0] += 1; alvo[1] += r.metros; alvo[2] += r.kg
    avisos = []
    mal = _seguro('Malharia', malharia_por_dia, avisos)
    prod_m = sum(v['total']['metros'] for k, v in mal.items() if ini.isoformat() <= k < fim.isoformat())
    pa, pm = _mes_anterior(ano, mes)
    ini_a, fim_a = date(pa, pm, 1), ini
    ate = min((date.today() - ini).days + 1, 31) if (ano, mes) == (date.today().year, date.today().month) else 31
    ra = [r for r in IndReprovaCru.query.filter(IndReprovaCru.data >= ini_a, IndReprovaCru.data < fim_a).all() if r.data.day <= ate]
    prod_a = sum(v['total']['metros'] for k, v in mal.items() if ini_a.isoformat() <= k < fim_a.isoformat() and int(k[8:10]) <= ate)
    por_maq = defaultdict(lambda: [0, 0.0])
    por_turno = [[0, 0.0] for _ in range(3)]
    for r in rs:
        por_maq[r.maquina or '—'][0] += 1
        por_maq[r.maquina or '—'][1] += r.metros
        if r.turno:
            por_turno[r.turno - 1][0] += 1
            por_turno[r.turno - 1][1] += r.metros
    return dict(ano=ano, mes=mes, pecas=len(rs), metros=round(sum(r.metros for r in rs), 1), kg=round(sum(r.kg for r in rs), 1),
                producao_metros=round(prod_m, 1), anterior=dict(pecas=len(ra), metros=round(sum(r.metros for r in ra), 1), producao_metros=round(prod_a, 1)),
                maquinas=sorted(({'maquina': k, 'pecas': v[0], 'metros': round(v[1], 1)} for k, v in por_maq.items()), key=lambda x: -x['metros']),
                turnos=[{'pecas': t[0], 'metros': round(t[1], 1)} for t in por_turno], avisos=avisos,
                defeitos=sorted(({'defeito': k, 'pecas': v[0], 'metros': round(v[1], 1), 'kg': round(v[2], 1)} for k, v in por_def.items()), key=lambda x: -x['metros']),
                dias=[{'data': k, 'pecas': v[0], 'metros': round(v[1], 1), 'kg': round(v[2], 1)} for k, v in sorted(por_dia.items())],
                linhas=[_nc_dict(r) for r in rs], erro=_ult_nc['erro'])


# ---------------------------------------------------------------- setups (acompanhamento de producao)
SETUPS_URL = os.getenv('SETUPS_URL', 'http://108.165.179.201:8090/api/v1').rstrip('/')
_CAMPOS_INICIO = ('inicio', 'data_inicio', 'dataInicio', 'iniciado_em', 'iniciadoEm', 'started_at', 'startedAt', 'start', 'data_hora', 'created_at', 'createdAt', 'data')


def _pin() -> str:
    return (os.getenv('SETUPS_TOKEN') or os.getenv('SETUPS_PIN') or '').strip()


def _get_setups(caminho: str, params: dict, tentativas: list = None):
    """GET na API de setups: autenticacao "Authorization: Bearer <token>" (descrita em /api/v1)."""
    import requests
    h = {'Accept': 'application/json'}
    if _pin():
        h['Authorization'] = f'Bearer {_pin()}'
    return requests.get(f'{SETUPS_URL}{caminho}', params=params, headers=h, timeout=(5, 25))      # 5 s para conectar; a resposta de /setups pode levar alguns segundos


def _lista(j) -> list:
    """Lista de setups da resposta: a propria lista, ou o 1o campo que seja uma lista (setups, items, data...)."""
    if isinstance(j, list):
        return j
    if isinstance(j, dict):
        for k in ('setups', 'items', 'data', 'results', 'dados', 'registros', 'lista'):
            if isinstance(j.get(k), list):
                return j[k]
        for v in j.values():
            if isinstance(v, list) and (not v or isinstance(v[0], dict)):
                return v
            if isinstance(v, dict):
                achou = _lista(v)
                if achou:
                    return achou
    return []


def _quando(item: dict):
    """Inicio do setup como datetime do horario de Brasilia (se vier com fuso, converte; se vier so a data, devolve a data)."""
    for k in _CAMPOS_INICIO:
        v = item.get(k)
        if not v:
            continue
        txt = str(v).strip().replace('Z', '+00:00')
        try:
            h = datetime.fromisoformat(txt)
        except ValueError:
            try:
                h = datetime.strptime(txt[:19], '%d/%m/%Y %H:%M:%S') if len(txt) > 10 else datetime.strptime(txt[:10], '%d/%m/%Y')
            except ValueError:
                continue
        if h.tzinfo is not None:
            h = (h - h.utcoffset()).replace(tzinfo=None) - timedelta(hours=3)
        return h, len(txt) > 10
    return None, False


def _setups_calc(ch: str) -> dict:
    """Setups do dia de producao `ch`. A API (/api/v1/setups) ja filtra de/ate pelo dia de producao (06h20 -> 06h19) e traz
    `diaProducao` e `turnoInicio` ("1º turno 20/10") de cada setup; sem esses campos o turno sai da hora de `inicio`."""
    import re
    r = _get_setups('/setups', {'de': ch, 'ate': ch, 'tipo': 'SETUP'})
    if r.status_code != 200:
        raise RuntimeError(f'HTTP {r.status_code} {r.text[:120]}')
    n, maqs, turnos, sem_turno = 0, set(), [0, 0, 0], 0
    for it in _lista(r.json()):
        tipo = str(it.get('tipo') or it.get('type') or 'SETUP').strip().upper()
        if tipo not in ('SETUP', ''):
            continue
        if it.get('diaProducao') and str(it['diaProducao'])[:10] != ch:
            continue
        n += 1
        m = re.match(r'\s*([123])', str(it.get('turnoInicio') or ''))
        if m:
            turnos[int(m.group(1)) - 1] += 1
        else:
            h, tem_hora = _quando(it)
            if h is not None and tem_hora:
                turnos[turno(h)] += 1
            else:
                sem_turno += 1
        maqs.add(str(it.get('maquina') or it.get('machine') or it.get('maquina_id') or ''))
    return dict(n=n, turnos=turnos if n and not sem_turno else None, maquinas=len(maqs - {''}))


def _setups_resumo(r) -> dict:
    r = r or {}
    return dict(setups=r.get('n'), setups_turnos=r.get('turnos'))


def setups_do_dia(dia: date):
    return _swr('setups', 120, _setups_calc, dia.isoformat())


def setups_diagnostico() -> dict:
    import requests
    out = dict(url=SETUPS_URL, token_configurado=bool(_pin()))
    try:
        r0 = requests.get(SETUPS_URL, headers={'Accept': 'application/json'}, timeout=15)
        out['descricao'] = dict(status=r0.status_code, corpo=r0.text[:1500])
        hoje = date.today()
        r = _get_setups('/setups', {'de': (hoje - timedelta(days=2)).isoformat(), 'ate': hoje.isoformat()})
        out['setups'] = dict(status=r.status_code)
        if r.status_code == 200:
            j = r.json()
            itens = _lista(j)
            out['setups'].update(total=len(itens), chaves=sorted(itens[0].keys()) if itens and isinstance(itens[0], dict) else None, exemplos=itens[:2], corpo_bruto=None if itens else r.text[:600], parametros=dict(de=(hoje - timedelta(days=2)).isoformat(), ate=hoje.isoformat()))
        else:
            out['setups']['corpo'] = r.text[:300]
    except Exception as e:  # noqa: BLE001
        out['erro'] = f'{type(e).__name__}: {str(e)[:200]}'
    return out


# ---------------------------------------------------------------- estoque de acabados (relatorio 040 = saldo 1 dos lotes)
BATCH_URL = f'{BASE_URL}/api/totvsmoda/product/v2/batch/search'
_est = {'ts': None, 'erro': None}


def _unidades(cli, codigos: list) -> None:
    faltam = [c for c in codigos if not db.session.get(IndProdutoUn, c)]
    for i in range(0, len(faltam), 50):
        lote = faltam[i:i + 50]
        r = cli.post(PROD_URL, dict(filter=dict(productCodeList=lote), option=dict(branchInfoCode=EMPRESA), page=1, pageSize=50))
        if r.status_code != 200:
            raise RuntimeError(f'produtos: HTTP {r.status_code} {r.text[:160]}')
        achados = {p['productCode']: p for p in r.json().get('items') or []}
        for c in lote:
            p = achados.get(c) or {}
            db.session.add(IndProdutoUn(codigo=c, nome=(p.get('productName') or '')[:200], unidade=str(p.get('measuredUnit') or '').strip().upper()[:5],
                                        acabado=bool(p.get('isFinishedProduct')) and not p.get('isRawMaterial')))
        db.session.commit()


def atualizar_estoque(app) -> bool:
    """Foto do estoque de acabados (saldo 1 dos lotes, so tecido acabado em metros): o que aparece e entrada, o que some e saida."""
    from app.relatorio_fat.agenda import agora_brasilia
    if not _trava.acquire(blocking=False):
        return False
    try:
        with app.app_context():
            cli = api.TotvsModaClient()
            cli.autenticar()
            atuais, pag = {}, 1
            while True:
                r = cli.post(BATCH_URL, dict(filter=dict(branchCode=EMPRESA, balanceCodeList=[1], hasBalance=True), page=pag, pageSize=1000))
                if r.status_code != 200:
                    raise RuntimeError(f'lotes: HTTP {r.status_code} {r.text[:160]}')
                j = r.json()
                for lote in j.get('items') or []:
                    for bi in lote.get('batchItem') or []:
                        bal = _num(bi.get('balance')) or 0
                        if bal > 0:
                            atuais[(lote['batchCode'], lote['productCode'], bi.get('batchItemSequential'))] = bal
                if not j.get('hasNext') or pag >= 30:
                    break
                pag += 1
            _unidades(cli, sorted({k[1] for k in atuais}))
            ok = {c for (c,) in db.session.query(IndProdutoUn.codigo).filter(IndProdutoUn.acabado.is_(True), IndProdutoUn.unidade == 'MT').all()}
            atuais = {k: v for k, v in atuais.items() if k[1] in ok}
            agora = agora_brasilia()
            primeira = IndEstoqueItem.query.count() == 0
            todos = {(x.lote, x.produto, x.seq): x for x in IndEstoqueItem.query.all()}
            for k, bal in atuais.items():
                x = todos.get(k)
                if x is None:
                    db.session.add(IndEstoqueItem(lote=k[0], produto=k[1], seq=k[2], metros=bal, visto_em=agora, baseline=primeira))
                elif x.saiu_em is not None:                         # a peca voltou ao estoque (estorno, nova entrada): conta como entrada de novo
                    x.saiu_em, x.visto_em, x.baseline, x.metros = None, agora, False, bal
                else:
                    x.metros = bal
            for k, x in todos.items():
                if x.saiu_em is None and k not in atuais:
                    x.saiu_em = agora
            db.session.commit()
            _est.update(ts=agora.isoformat(timespec='seconds'), erro=None)
        return True
    except Exception as e:  # noqa: BLE001
        traceback.print_exc()
        try:
            db.session.rollback()
        except Exception:  # noqa: BLE001
            pass
        _est['erro'] = str(e)[:200]
        return False
    finally:
        _trava.release()


def estoque_dia(dia: date) -> dict:
    """Entradas e saidas do dia de producao no estoque de acabados (por turno) e o saldo atual, vistos pela foto a cada 2 min."""
    ini = datetime.combine(dia, datetime.min.time()) + INICIO_DIA
    fim = ini + timedelta(days=1)
    desde = db.session.query(db.func.min(IndEstoqueItem.visto_em)).scalar()
    ent = [dict(pecas=0, metros=0.0) for _ in range(3)]
    sai = dict(pecas=0, metros=0.0)
    for x in IndEstoqueItem.query.filter(IndEstoqueItem.visto_em >= ini, IndEstoqueItem.visto_em < fim, IndEstoqueItem.baseline.is_(False)).all():
        t = ent[turno(x.visto_em)]
        t['pecas'] += 1
        t['metros'] += x.metros or 0
    for x in IndEstoqueItem.query.filter(IndEstoqueItem.saiu_em >= ini, IndEstoqueItem.saiu_em < fim).all():
        sai['pecas'] += 1
        sai['metros'] += x.metros or 0
    saldo = db.session.query(db.func.count(IndEstoqueItem.id), db.func.coalesce(db.func.sum(IndEstoqueItem.metros), 0)).filter(IndEstoqueItem.saiu_em.is_(None)).first()
    for t in ent:
        t['metros'] = round(t['metros'], 1)
    return dict(monitorando_desde=desde.isoformat(timespec='minutes') if desde else None, cobre_o_dia=bool(desde) and desde <= ini,
                entradas=ent, entrada_total=dict(pecas=sum(t['pecas'] for t in ent), metros=round(sum(t['metros'] for t in ent), 1)),
                saidas=dict(pecas=sai['pecas'], metros=round(sai['metros'], 1)), saldo=dict(pecas=int(saldo[0] or 0), metros=round(float(saldo[1] or 0), 1)),
                atualizado=_est['ts'], erro=_est['erro'])


def _scrap_calc(ch: str) -> dict:
    from app.indicadores import scrap
    dia = date.fromisoformat(ch)
    vivo = defaultdict(float)
    for r in IndReprovaCru.query.filter(IndReprovaCru.data >= date(dia.year, 1, 1), IndReprovaCru.data <= dia).all():
        vivo[r.data] += r.kg or 0
    return scrap.scrap_dia(dia, dict(vivo))


def scrap_do_dia(dia: date):
    return _swr('scrap', 120, _scrap_calc, dia.isoformat())


def _scrap_ou_none(avisos: list, dia: date):
    """Scrap do dia; se ainda nao ha planilha, devolve None (a tela mostra o botao para enviar), sem aviso de erro."""
    try:
        return scrap_do_dia(dia)
    except RuntimeError as e:
        if 'SEM_PLANILHA' in str(e):
            return None
        avisos.append(f'Scrap: {e}')
        return None
    except Exception as e:  # noqa: BLE001
        traceback.print_exc()
        avisos.append(f'Scrap: {type(e).__name__}: {str(e)[:160]}')
        return None


def diagnostico_dia(dia: date) -> dict:
    """Conferencia: de onde vem cada numero do dia (malharia, acabamento, faturamento), com as primeiras pecas e as contas."""
    from app.main.routes import _api_baixas, _hub_movimentos_json
    ch = dia.isoformat()
    linhas = json.loads('[' + _api_baixas(dict(_hub_movimentos_json())) + ']')
    mal = []
    for r in linhas:
        if r.get('cod_artigo') and 'amostra' in r['cod_artigo'].lower():
            continue
        h = datetime.fromisoformat(r['dt_real'].replace('Z', '')) - timedelta(hours=3)
        if dia_producao(h).isoformat() == ch:
            mal.append(dict(op=r.get('op'), artigo=(r.get('cod_artigo') or '')[:50], hora=h.strftime('%H:%M'), turno=turno(h) + 1, metros_reais=r.get('qt_movimento'), metro_padrao=r.get('metro_padrao'), kg=r.get('qt_pesos')))
    aca = []
    prod = {p.codigo: p for p in IndProduto.query.all()}
    for r in IndAcabamento.query.order_by(IndAcabamento.mudou).all():
        if dia_producao(r.mudou).isoformat() == ch:
            p = prod.get(r.produto)
            aca.append(dict(transacao=r.codigo, operacao=r.operacao, hora=r.mudou.strftime('%H:%M'), turno=turno(r.mudou) + 1, produto=r.produto, nome=(p.nome if p else '')[:40], metros=r.metros, gramatura=p.gramatura if p else None, largura=p.largura if p else None, kg=r.kg))
    soma = lambda l, k: round(sum((x.get(k) or 0) for x in l), 1)
    return dict(dia=ch, regra_do_dia='dia de producao = 06:20 ate 06:20 do dia seguinte; turnos 06:20-14:20, 14:20-22:20, 22:20-06:20',
                malharia=dict(pecas=len(mal), metros_reais=soma(mal, 'metros_reais'), metros_padrao=soma(mal, 'metro_padrao'), kg=soma(mal, 'kg'), pecas_sem_metros=sum(1 for x in mal if not (x.get('metros_reais') or x.get('metro_padrao'))), pecas_sem_op=sum(1 for x in mal if not x.get('op')), pecas_sem_kg=sum(1 for x in mal if not x.get('kg')), primeiras=mal[:8], sem_metros_exemplos=[x for x in mal if not (x.get('metros_reais') or x.get('metro_padrao'))][:6]),
                acabamento=dict(pecas=len(aca), metros=soma(aca, 'metros'), kg=soma(aca, 'kg'), por_operacao={str(o): sum(1 for x in aca if x['operacao'] == o) for o in {x['operacao'] for x in aca}}, sem_gramatura_ou_largura=sum(1 for x in aca if not x['kg']), primeiras=aca[:8]),
                faturamento_hoje_nfs={k: v for k, v in __import__('app.totvs.expedicao', fromlist=['x']).vendas_por_dia(dia.year, dia.month).items() if k == ch})


def comparar_planilha(dia: date) -> dict:
    """Compara, maquina a maquina e turno a turno, o que o Gestao calcula com a planilha INDICADORES (abas MALHARIA e ACABAMENTO) para o dia."""
    import json
    import openpyxl
    from pathlib import Path
    from app.indicadores import scrap
    from app.models import BaixaFio
    plan, aca_plan = {}, None
    arq = scrap.caminho()
    if arq and arq.exists():
        wb = openpyxl.load_workbook(arq, read_only=True, data_only=True)
        for r in wb['MALHARIA'].iter_rows(min_row=2, max_col=40, values_only=True):
            d = r[1].date() if isinstance(r[1], datetime) else None
            if d == dia and r[0] is not None:
                plan[int(r[0])] = dict(metros=[r[17] or 0, r[18] or 0, r[19] or 0], pecas=[r[33] or 0, r[34] or 0, r[35] or 0], kg=[round(r[37] or 0, 1), round(r[38] or 0, 1), round(r[39] or 0, 1)])
        for r in wb['ACABAMENTO'].iter_rows(min_row=2, max_col=13, values_only=True):
            if isinstance(r[0], datetime) and r[0].date() == dia:
                aca_plan = dict(pecas=[r[2], r[3], r[4]], metros=[round(r[6] or 0, 1), round(r[7] or 0, 1), round(r[8] or 0, 1)], kg=[round(r[10] or 0, 1), round(r[11] or 0, 1), round(r[12] or 0, 1)])
        wb.close()
    if not plan:                                    # sem a planilha no servidor: usa o retrato do dia guardado no sistema (26/10/2026, enviado pela diretoria)
        f = Path(__file__).with_name('planilha_verdade.json')
        v = json.loads(f.read_text(encoding='utf-8')).get(dia.isoformat()) if f.exists() else None
        if not v:
            return dict(erro='Planilha INDICADORES nao encontrada no servidor (use "enviar planilha atualizada" na aba Indicadores).')
        plan = {int(k): x for k, x in v['malharia'].items()}
        aca_plan = v.get('acabamento')
    nos, pecas_por_maq = {}, {}
    for b in BaixaFio.query.order_by(BaixaFio.mudou).all():
        if not b.mudou or dia_producao(b.mudou) != dia or (b.artigo_desc and 'amostra' in b.artigo_desc.lower()):
            continue
        t = turno(b.mudou)
        m = int(b.maquina) if b.maquina is not None else -1
        x = nos.setdefault(m, dict(metros=[0.0] * 3, pecas=[0] * 3, kg=[0.0] * 3))
        padrao = b.metro_padrao or b.metros or 0
        x['pecas'][t] += 1
        x['metros'][t] += padrao
        x['kg'][t] += _peso_ind(b, padrao)
        pecas_por_maq.setdefault(m, []).append(dict(op=b.op, artigo=(b.artigo_desc or '')[:40], hora=b.mudou.strftime('%H:%M'), turno=t + 1, metro_padrao=b.metro_padrao, metros_reais=b.metros, kg_baixa_fio=b.kg, gramatura=b.gramatura, largura=b.largura))
    difere = []
    for m in sorted(set(plan) | set(nos)):
        p_, n_ = plan.get(m), nos.get(m)
        pm = p_['metros'] if p_ else [0, 0, 0]
        nm = [round(v, 1) for v in n_['metros']] if n_ else [0, 0, 0]
        pk = p_['kg'] if p_ else [0, 0, 0]
        nk = [round(v, 1) for v in n_['kg']] if n_ else [0, 0, 0]
        pp = p_['pecas'] if p_ else [0, 0, 0]
        npp = n_['pecas'] if n_ else [0, 0, 0]
        if any(abs(a - b) > 0.5 for a, b in zip(pm, nm)) or any(abs(a - b) > 1 for a, b in zip(pk, nk)) or pp != npp:
            difere.append(dict(maquina=m, planilha=dict(metros=pm, pecas=pp, kg=pk), gestao=dict(metros=nm, pecas=npp, kg=nk), pecas_gestao=pecas_por_maq.get(m, [])[:14]))
    tot = lambda d, k: round(sum(sum(v[k]) for v in d.values()), 1)
    return dict(dia=dia.isoformat(), planilha_malharia=dict(metros=tot(plan, 'metros'), pecas=tot(plan, 'pecas'), kg=tot(plan, 'kg')), gestao_malharia=dict(metros=tot(nos, 'metros'), pecas=tot(nos, 'pecas'), kg=tot(nos, 'kg')),
                maquinas_que_diferem=difere, acabamento_planilha=aca_plan, acabamento_gestao=(acabamento_por_dia().get(dia.isoformat()) or {}).get('t'))


def _peso_ind(b, metros: float) -> float:
    """kg da peca: metros x gramatura x largura / 1000 (campos do produto); senao pelo nome do artigo; senao so o fio da baixa."""
    import re
    if b.gramatura and b.largura and metros:
        return round(metros * b.gramatura * b.largura / 1000, 2)
    g = re.search(r'(\d{2,3})\s*G\b', b.artigo_desc or '', re.I)
    l = re.search(r'(\d+[.,]\d+)\s*L\b', b.artigo_desc or '', re.I)
    if g and l and metros:
        return round(metros * int(g.group(1)) * float(l.group(1).replace(',', '.')) / 1000, 2)
    return round(b.kg or 0, 2)
