"""Scrap lido da planilha INDICADORES (Excel), como no relatorio da reuniao diaria.

A planilha vem do botao "Enviar planilha" da aba Indicadores (fica em instance/indicadores) ou do .env: INDICADORES_XLSX=C:\\caminho\\INDICADORES_2026.xlsm (o arquivo e so lido, nunca alterado).
Abas usadas: MALHARIA (kg produzido, col. T. KG), RETALHO (kg por local), REPROVA CRU e REPROVA ACA. (kg reprovado).
Scrap do dia = (retalho + reprova cru + reprova acabamento) / kg produzido no dia. Semana = bloco de 7 dias a partir de 1o de janeiro (como a planilha),
mes e ano acumulados. A planilha precisa estar salva com os dados do dia para ele aparecer.
"""
from __future__ import annotations

import os
import threading
import warnings
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

META = 0.04
_cache = {'chave': None, 'dados': None}
_trava = threading.Lock()


def pasta_planilha() -> Path:
    from flask import current_app
    return Path(current_app.instance_path) / 'indicadores'


def caminho():
    """Planilha INDICADORES: a do .env (INDICADORES_XLSX) ou a enviada pelo botao "Enviar planilha" da aba (instance/indicadores)."""
    c = os.getenv('INDICADORES_XLSX', '').strip().strip('"')
    if c:
        return Path(c)
    try:
        achados = sorted(pasta_planilha().glob('INDICADORES.xls*'), key=lambda x: x.stat().st_mtime, reverse=True)
    except Exception:  # noqa: BLE001
        achados = []
    return achados[0] if achados else None


def semana_planilha(d: date) -> int:
    """Numero da semana como na planilha: blocos de 7 dias a partir de 1o de janeiro (19/10 = semana 42)."""
    return (d.timetuple().tm_yday - 1) // 7 + 1


def _dia(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0


def _ler(arq: Path) -> dict:
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        wb = openpyxl.load_workbook(arq, read_only=True, data_only=True)
        prod, ret, rc, ra = defaultdict(float), defaultdict(lambda: defaultdict(float)), defaultdict(float), defaultdict(float)
        for r in wb['MALHARIA'].iter_rows(min_row=2, max_col=37, values_only=True):
            d = _dia(r[1])
            if d:
                prod[d] += _num(r[36])
        for r in wb['RETALHO'].iter_rows(min_row=2, max_col=4, values_only=True):
            d = _dia(r[0])
            if d:
                ret[d][str(r[2] or 'OUTROS').strip().upper()] += _num(r[1])
        for r in wb['REPROVA CRU'].iter_rows(min_row=2, max_col=6, values_only=True):
            d = _dia(r[0])
            if d:
                rc[d] += _num(r[5])
        for r in wb['REPROVA ACA.'].iter_rows(min_row=2, max_col=5, values_only=True):
            d = _dia(r[0])
            if d:
                ra[d] += _num(r[4])
        wb.close()
    return dict(prod=prod, ret=ret, rc=rc, ra=ra)


def _dados():
    arq = caminho()
    if not arq:
        raise RuntimeError('SEM_PLANILHA')
    if not arq.exists():
        raise RuntimeError(f'Planilha não encontrada: {arq}')
    chave = (str(arq), arq.stat().st_mtime)
    with _trava:
        if _cache['chave'] != chave:
            _cache.update(chave=chave, dados=_ler(arq))
    return _cache['dados'], datetime.fromtimestamp(arq.stat().st_mtime)


def _soma(dados, dias, rc_extra=None):
    ret = sum(sum(dados['ret'].get(d, {}).values()) for d in dias)
    ultimo = max(dados['rc']) if dados['rc'] else date.min                       # depois da ultima data da planilha, vale a NC viva do SharePoint
    rc = sum((rc_extra or {}).get(d, 0.0) if d > ultimo else dados['rc'].get(d, 0.0) for d in dias)
    ra = sum(dados['ra'].get(d, 0.0) for d in dias)
    prod = sum(dados['prod'].get(d, 0.0) for d in dias)
    return ret, rc, ra, prod


def scrap_dia(dia: date, rc_vivo: dict | None = None) -> dict:
    """rc_vivo: {data: kg} da reprova cru das NCs do SharePoint (mais atual que a planilha); se faltar, usa a da planilha."""
    dados, salvo = _dados()
    ini_sem = date(dia.year, 1, 1) + timedelta(days=7 * ((dia.timetuple().tm_yday - 1) // 7))
    sem = [ini_sem + timedelta(days=i) for i in range(7) if ini_sem + timedelta(days=i) <= dia]
    mes = [date(dia.year, dia.month, d) for d in range(1, dia.day + 1)]
    ano = [date(dia.year, 1, 1) + timedelta(days=i) for i in range((dia - date(dia.year, 1, 1)).days + 1)]

    def pct(dias):
        r, c, a, p = _soma(dados, dias, rc_vivo)
        return (r + c + a) / p if p else None
    ret, rc, ra, prod = _soma(dados, [dia], rc_vivo)
    locais = dados['ret'].get(dia, {})
    return dict(retalho_kg=round(ret, 2), retalho_locais={k: round(v, 2) for k, v in sorted(locais.items(), key=lambda x: -x[1])},
                reprova_cru_kg=round(rc, 2), reprova_aca_kg=round(ra, 2), producao_kg=round(prod, 1),
                scrap_kg=round(ret + rc + ra, 2), dia=pct([dia]), semana=pct(sem), mes=pct(mes), ano=pct(ano), meta=META,
                planilha_salva_em=salvo.isoformat(timespec='minutes'), tem_dados_do_dia=prod > 0 or ret > 0 or ra > 0)
